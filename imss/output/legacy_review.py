"""旧版 13 列 Excel/CSV 与候选确认的只读兼容层；新文件统一使用八列。"""
import csv
from pathlib import Path
from openpyxl import load_workbook
from imss.core.models import KINDS, Material, money
from imss.output.candidate_review import selected_candidates
from imss.core.candidates import pairing_candidates
CSV_FIELDS = ['source', 'kind', 'item', 'amount', 'group', 'title', 'confirmed', 'issues']

COLUMNS = [('severity', '问题程度'), ('issues', '问题与处理建议'), ('confirmed', '已核对'),
           ('item', '物品名称'), ('amount', '金额'), ('kind', '材料类型'),
           ('group', '分组'), ('source', '原文件（点击打开）'), ('matched', '配对原件'), ('title', '标题')]

OPTIONAL_COLUMNS = [('merchant', '商家'), ('transaction_date', '交易或开票日期'), ('evidence', '字段证据（页码与坐标）')]



def read_review(path: Path):
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream)
            if not set(CSV_FIELDS).issubset(reader.fieldnames or []):
                raise ValueError('核对清单缺少必需列')
            return list(reader)
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        sheet = workbook['核对清单']
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=4, max_row=4))]
        labels = {label: key for key, label in COLUMNS + OPTIONAL_COLUMNS}
        if len(headers) != len(set(headers)) or not {label for _, label in COLUMNS}.issubset(headers):
            raise ValueError('核对清单缺少必需列或列名重复')
        result = []
        for cells in sheet.iter_rows(min_row=5):
            if not any(cell.value is not None for cell in cells):
                continue
            if any(cell.data_type == 'f' for cell in cells):
                raise ValueError('核对清单不支持公式，请直接填写文本或金额')
            result.append({labels[header]: '' if cell.value is None else str(cell.value)
                           for header, cell in zip(headers, cells) if header in labels})
        return result
    finally:
        workbook.close()



def apply_review(path: Path, materials: list[Material]):
    """只应用人工确认行；重算配对仍会阻止缺件、金额冲突或多发票。"""
    from imss.output.review import review_rows
    by_source = {m.source: m for m in materials}
    baseline = {row['source']: row for row in review_rows(materials)}
    proposals = selected_candidates(path, pairing_candidates(materials))
    edited = set()
    seen = set()
    for row in read_review(path):
        source = row['source']
        if source not in by_source or source in seen:
            raise ValueError(f'核对清单有未知或重复 source: {source}')
        seen.add(source)
        kind = row['kind'].strip()
        if kind not in KINDS:
            raise ValueError(f'材料类型无效: {kind}')
        if row['confirmed'].strip().lower() not in ('1', '1.0', 'yes', 'true', '是'):
            def comparable(field, value):
                value = value.strip()
                if field == 'amount' and value:
                    try:
                        return money(value)
                    except ValueError:
                        pass
                return value
            if any(comparable(field, row[field]) != comparable(field, baseline[source][field])
                   for field in ('kind', 'item', 'amount', 'group', 'title')):
                edited.add(source)
                by_source[source].issues.append('核对表修改尚未确认，请核对后将已核对填1')
            continue
        m = by_source[source]
        m.kind, m.item = kind, row['item'].strip()
        m.amount = money(row['amount']) if row['amount'].strip() else ''
        m.group, m.title = row['group'].strip(), row['title'].strip()
        m.confirmed = True
        m.match_basis = 'reviewed'
        m.issues.clear()
        if not m.extension:
            m.issues.append('不支持的文件格式，不能通过核对清单改为可处理格式')
    if seen != set(by_source):
        raise ValueError('核对清单不可删除行；无需处理的材料请将类型改为 未知')
    for proposal in proposals:
        members = [by_source[source] for source in proposal['members']]
        if any(m.confirmed or m.source in edited for m in members):
            raise ValueError('同一原件不能同时在两张表修改或确认；请保留一种确认方式')
        for material in members:
            material.group = 'candidate-' + proposal['id']
            if material.kind != '购买清单':
                material.item = proposal['item']
            material.amount = material.amount or proposal['amount']
            material.confirmed = True
            material.match_basis = 'candidate_reviewed'
            material.issues.clear()
            if not material.extension:
                material.issues.append('不支持的文件格式，不能通过候选确认改为可处理格式')

