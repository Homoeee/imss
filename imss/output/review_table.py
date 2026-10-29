"""八列核对表的写出、人工覆盖与状态读取；原始识别数据始终留在 JSON。"""
import csv
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from pathlib import Path
import math
import re
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from imss.core.models import KINDS, SPECIALS, STANDALONE_SHEETS, money, inside

COLUMNS = [('source', '原发票完整名称'), ('issues', '问题'), ('ocr', '文字识别结果'),
           ('advice', '处理建议'), ('kind', '类型'), ('item', '名称'), ('amount', '金额'),
           ('confirmed', '校验标志位')]
COLORS = {'严重': 'FADBD8', '待确认': 'FCE4C5', '提示': 'FFF2CC', '正常': 'E2F0D9'}
WIDTHS = (38, 42, 62, 42, 15, 25, 13, 16)
VERIFIED, WAITING = '已校验', '待校验'


def status(value):
    value = str(value or '').strip()
    if value not in ('', VERIFIED, WAITING):
        raise ValueError('校验标志位只能选择“已校验”或“待校验”')
    return value or WAITING


def is_simple_review(path):
    if path.suffix.lower() == '.csv':
        return False
    book = load_workbook(path, read_only=True)
    try:
        headers = [c.value for c in book['核对清单'][4]]
        return '校验标志位' in headers
    finally:
        book.close()


def read_review(path):
    from imss.output.legacy_review import read_review as read_legacy
    if not is_simple_review(path):
        return [dict(row, _format='legacy') for row in read_legacy(path)]
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book['核对清单']
        headers = [c.value for c in sheet[4]]
        expected = [label for _, label in COLUMNS]
        if headers != expected:
            raise ValueError('核对清单必须保留八列原始列名及顺序')
        rows = []
        for cells in sheet.iter_rows(min_row=5):
            if not any(c.value is not None for c in cells):
                continue
            if any(c.data_type == 'f' for c in cells):
                raise ValueError('核对清单不支持公式，请直接填写文本或金额')
            row = {key: '' if cell.value is None else str(cell.value)
                   for (key, _), cell in zip(COLUMNS, cells)}
            row['confirmed'] = status(row['confirmed'])
            row['_format'] = 'simple'
            row['severity'] = '正常' if not row['issues'] else '严重'
            rows.append(row)
        return rows
    finally:
        book.close()


def apply_review(path, materials):
    """新表只覆盖非空字段；每次重新选择全部已校验材料，清除旧人工分组。"""
    from imss.output.legacy_review import apply_review as apply_legacy
    if not is_simple_review(path):
        apply_legacy(path, materials)
        return 'legacy'
    rows = read_review(path)
    by_source = {m.source: m for m in materials}
    seen = set()
    for row in rows:
        source = row['source']
        if source not in by_source or source in seen:
            raise ValueError(f'核对清单有未知或重复原文件名称: {source}')
        seen.add(source)
        m = by_source[source]
        m.confirmed = row['confirmed'] == VERIFIED
        m.group = ''
        m.match_basis = 'reviewed' if m.confirmed else ''
        if not m.confirmed:
            continue
        if row['kind'].strip():
            if row['kind'].strip() not in KINDS:
                raise ValueError(f'材料类型无效: {source}：{row["kind"]}')
            m.kind = row['kind'].strip()
        if row['item'].strip():
            m.item = row['item'].strip()
            if m.kind in (*SPECIALS, *STANDALONE_SHEETS):
                m.title = m.item
        if row['amount'].strip():
            m.amount = money(row['amount'])
        m.issues.clear()
        if not m.extension:
            m.issues.append('不支持的文件格式，不能通过校验标志位改为可处理格式')
        if m.kind in ('行程单', '运单明细'):
            m.category = '行程' if m.kind == '行程单' else '快递'
    if seen != set(by_source):
        raise ValueError('核对清单不可删除行；不处理的材料请选择待校验')
    return 'simple'


def recognized_text(material):
    prefix = (f'识别类型：{material.kind}\n识别名称：{material.item or "未识别"}\n'
              f'识别金额：{material.amount or "未识别"}\n\nOCR原文：\n')
    full = prefix + material.text
    cleaned = excel_text(full)
    if cleaned != full:
        full = cleaned + '\n（已过滤Excel不支持的控制字符，完整原文保存在核对清单.json。）'
    if len(full) > 32767:
        full = full[:32700] + '\n（Excel单元格长度受限，完整原文保存在同目录核对清单.json。）'
    return full


def excel_text(value):
    # OCR/PDF 可能含 U+FFFE、孤立代理项等 XML 禁止字符，原始结果仍保存在 JSON。
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', str(value))


def advice_for(row):
    issues = row.get('issues', '')
    if not issues:
        return '识别及配对无问题，默认已校验；如需暂不生成，请改为待校验。'
    actions = []
    if any(word in issues for word in ('物品', '名称', '多个商品')):
        actions.append('对照原件，在“名称”填写正确名称；同一笔附件使用一致名称。')
    if any(word in issues for word in ('金额', '总额', '小数点', '实付')):
        actions.append('核查实际总额，在“金额”填写正确数值；不要填写单价或税额。')
    if any(word in issues for word in ('类型', '未知')):
        actions.append('从“类型”下拉列表选择正确材料类型。')
    if any(word in issues for word in ('配对', '匹配', '分组', '候选')):
        actions.append('核对发票及附件的类型、名称和金额；build将重新配对，多个相同候选不会任意选择。')
    if '缺少' in issues:
        actions.append('检查对应附件是否仍为待校验；原件确实缺失时需补齐材料。')
    if not actions:
        actions.append('点击原文件对照文字识别结果，检查问题并按需修改。')
    actions.append('核实后选择“已校验”，保存并关闭Excel，再执行build。')
    return '\n'.join(actions)


def view_rows(materials, states=None, diagnostics=None):
    from imss.output.review import review_rows
    states = states or {}
    diagnostics = diagnostics or {r['source']: r for r in review_rows(materials)}
    result = []
    for material in materials:
        row = dict(diagnostics[material.source])
        state = states.get(material.source, {})
        row['ocr'] = recognized_text(material)
        row['advice'] = advice_for(row)
        row['issues'] = row['issues'].replace('请填写分组', '请核对类型、名称和金额后重新配对')
        row['kind'] = state.get('kind', '')
        row['item'] = state.get('item', '')
        row['amount'] = state.get('amount', '')
        row['confirmed'] = status(state['confirmed']) if 'confirmed' in state else (
            VERIFIED if not row['issues'] else WAITING)
        result.append(row)
    priorities = {'严重': 0, '待确认': 1, '提示': 2, '正常': 3}
    return sorted(result, key=lambda row: (priorities[row['severity']], row['source']))


def migrated_states(materials, rows):
    """旧表人工编辑和候选确认转成覆盖项；不把旧自动分组锁定在新表中。"""
    originals = {m.source: m for m in materials}
    states = {}
    for row in rows:
        if row.get('_format') == 'simple':
            states[row['source']] = row
            continue
        m = originals[row['source']]
        confirmed = str(row.get('confirmed', '')).strip().lower() in ('1', '1.0', '是', 'yes', 'true')
        def same(field):
            left, right = row.get(field, '').strip(), str(getattr(m, field)).strip()
            if field == 'amount' and left and right:
                try:
                    return money(left) == money(right)
                except ValueError:
                    pass
            return left == right
        changed = any(not same(field) for field in ('kind', 'item', 'amount', 'title'))
        states[m.source] = dict(confirmed=VERIFIED if confirmed else (VERIFIED if not row.get('issues') and not changed else WAITING),
            **{field: row.get(field, '') if not same(field) else ''
               for field in ('kind', 'item', 'amount')})
        if (row.get('title', '').strip() != m.title
                and row.get('kind', m.kind) in (*SPECIALS, *STANDALONE_SHEETS)):
            states[m.source]['item'] = row['title'].strip()
    return states


def write_review(path, materials, root=None, states=None):
    if path.suffix.lower() == '.csv':
        # CSV 只保留历史接口；正式核对工作流始终输出新版 Excel。
        from imss.output.legacy_review import CSV_FIELDS
        from imss.output.review import review_rows
        with path.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=CSV_FIELDS, extrasaction='ignore')
            writer.writeheader()
            writer.writerows(review_rows(materials))
        return
    write_review_rows(path, view_rows(materials, states), root)


def write_review_rows(path, rows, root, proposals=None):
    book = Workbook()
    sheet = book.active
    sheet.title = '核对清单'
    notes = [
        '核对清单：红色=严重问题；橙色=疑似配对；黄色=识别提示；绿色=识别及配对通过。先核查有问题的材料。',
        '点击原文件打开原件。类型、名称、金额留空沿用识别值；填写则覆盖。仅已校验材料参与build重配对。C列全文可在编辑栏展开查看；不要改文件名或删除行。',
        f'保存并关闭Excel后执行：python main.py build "{path.resolve()}"。校验标志位可选：已校验、待校验。']
    for index, text in enumerate(notes, 1):
        sheet.cell(index, 1, text)
        sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=8)
        sheet.cell(index, 1).alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[index].height = 30
    sheet.append([label for _, label in COLUMNS])
    for cell in sheet[4]:
        cell.fill = PatternFill('solid', fgColor='24445C')
        cell.font = Font(color='FFFFFF', bold=True, size=11)
    for index, original in enumerate(rows, 5):
        row = dict(original)
        target = (root / row['source']).resolve() if root else None
        if target and (not inside(target, root) or not target.is_file()):
            row['severity'] = '严重'
            row['issues'] = '；'.join(filter(None, [row['issues'], '原文件不存在或路径越界，链接未生成']))
            row['advice'] += '\n检查原始输入目录，恢复对应原件后刷新核对表。'
        for column, (key, _) in enumerate(COLUMNS, 1):
            cell = sheet.cell(index, column, excel_text(row.get(key, '')))
            cell.data_type = 's'
            cell.number_format = '@'
            if key == 'amount' and row.get(key):
                try:
                    amount = Decimal(str(row[key]))
                    if amount.is_finite():
                        cell.value = amount
                        cell.number_format = '0.00'
                except InvalidOperation:
                    pass
            cell.fill = PatternFill('solid', fgColor=COLORS[row['severity']])
            cell.font = Font(name='微软雅黑', size=11, color='1261A0' if column >= 5 else '202B33',
                             bold=key == 'issues' and row['severity'] != '正常')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if key == 'source' and target and target.is_file() and inside(target, root):
                # Excel 的本地文件链接使用 Windows 原生绝对路径，中文和空格不做 URL 编码。
                cell.hyperlink = str(target)
                cell.font = Font(name='微软雅黑', size=11, color='1261A0', underline='single')
        lines = max(max(1, str(row.get(key, '')).count('\n') + 1,
                        math.ceil(len(str(row.get(key, ''))) * 1.5 / width))
                    for (key, _), width in zip(COLUMNS, WIDTHS))
        # 先显示识别摘要与问题，避免一条长 OCR 占满屏幕；全文可在编辑栏展开。
        sheet.row_dimensions[index].height = min(126, max(60, lines * 15 + 8))
    for column, width in zip('ABCDEFGH', WIDTHS):
        sheet.column_dimensions[column].width = width
    for column, values in [('E', ','.join(KINDS)), ('H', VERIFIED + ',' + WAITING)]:
        validation = DataValidation(type='list', formula1='"' + values + '"', allow_blank=column == 'E')
        validation.showDropDown = False
        validation.showErrorMessage = True
        validation.errorStyle = 'stop'
        validation.error = '请选择下拉列表中的值。'
        sheet.add_data_validation(validation)
        validation.add(f'{column}5:{column}{max(5, sheet.max_row)}')
    sheet.freeze_panes = 'B5'  # 横向滚动填写时只固定原件名，避免窄屏被四个只读列占满。
    sheet.auto_filter.ref = f'A4:H{max(4, sheet.max_row)}'
    sheet.sheet_view.showGridLines = False
    temporary = path.with_suffix('.partial.xlsx')
    try:
        # 先完整序列化到内存，失败时不会留下占用中的半个 ZIP 文件。
        with io.BytesIO() as buffer:
            book.save(buffer)
            temporary.write_bytes(buffer.getvalue())
        book.close()
        temporary.replace(path)
    except PermissionError as exc:
        raise PermissionError(f'核对清单无法写入，请保存并关闭 Excel 后重试：{path}') from exc
    finally:
        book.close()
        temporary.unlink(missing_ok=True)


def refresh_review(path, materials, root, report, recognized=None):
    if path.suffix.lower() != '.xlsx':
        return
    from imss.output.review import review_rows
    recognized = recognized or materials
    previous = read_review(path)
    states = migrated_states(recognized, previous)
    if previous and previous[0].get('_format') == 'legacy':
        # 历史候选已应用的结果也要保留；写入的是覆盖项，不改原始识别字段。
        for m, original in zip(materials, recognized):
            if m.confirmed:
                states[m.source].update(confirmed=VERIFIED, **{
                    field: getattr(m, field) if getattr(m, field) != getattr(original, field) else states[m.source][field]
                    for field in ('kind', 'item', 'amount')})
    active = [m for m in materials if m.confirmed] if previous and previous[0].get('_format') == 'simple' else materials
    diagnostics = {row['source']: row for row in review_rows(deepcopy(recognized))}
    diagnostics.update({row['source']: row for row in review_rows(deepcopy(active))})
    by_group = {}
    for entry in report['pending']:
        sources = entry.get('sources') or ([entry['source']] if entry.get('source') else [])
        for source in sources:
            if source in diagnostics:
                row = diagnostics[source]
                row['issues'] = '；'.join(dict.fromkeys(filter(None, [row['issues'], *entry['issues']])))
                row['severity'] = '严重' if states[source]['confirmed'] == VERIFIED else row['severity']
    write_review_rows(path, view_rows(recognized, states, diagnostics), root)
