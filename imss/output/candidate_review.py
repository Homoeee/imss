"""候选批量确认表：确认的是整组证据，读取时重新计算并校验候选身份。"""
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

CANDIDATE_COLUMNS = [('confirmed', '整组已核对'), ('id', '候选编号'), ('status', '候选状态'),
                     ('rank', '订单内排名'), ('score', '证据排序分（非概率）'),
                     ('invoice', '发票原件'), ('order', '订单原件'), ('item', '发票物品名称'),
                     ('order_item', '订单原始名称'), ('amount', '发票金额'),
                     ('reasons', '配对依据'), ('risks', '需要核查的问题')]
YES = ('1', '1.0', 'yes', 'true', '是')


def write_candidates(workbook, proposals, root):
    sheet = workbook.create_sheet('候选配对')
    instructions = [
        '候选需要人工核查；排序分只是证据权重，不是识别准确率。名称、金额或编号冲突不能靠提高分数放行。',
        '点击发票和订单对照，核实整组所有字段和问题后，在 A 列填1；可筛选后批量粘贴1。同一订单、同一发票只能选一个候选。',
        '确认将统一为发票名称、补齐缺失金额并确认相关附件；需改金额或另行分组时请用“核对清单”，不要改本表其他列。保存关闭后执行 build。']
    for number, text in enumerate(instructions, 1):
        sheet.append([text])
        sheet.merge_cells(start_row=number, start_column=1, end_row=number, end_column=len(CANDIDATE_COLUMNS))
        sheet.cell(number, 1).alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[number].height = 32
    sheet.append([label for _, label in CANDIDATE_COLUMNS])
    for cell in sheet[4]:
        cell.fill = PatternFill('solid', fgColor='24445C')
        cell.font = Font(color='FFFFFF', bold=True)
    for index, proposal in enumerate(proposals, 5):
        color = 'E2F0D9' if str(proposal.get('confirmed', '')).lower() in YES else 'FCE4C5' if proposal['status'] == '双方唯一候选' else 'FADBD8'
        for column, (key, _) in enumerate(CANDIDATE_COLUMNS, 1):
            cell = sheet.cell(index, column, str(proposal.get(key, '')))
            cell.data_type = 's'
            if key in ('rank', 'score'):
                cell.value = int(proposal[key])
            cell.fill = PatternFill('solid', fgColor=color)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if key in ('invoice', 'order') and root:
                cell.hyperlink = (root / proposal[key]).resolve().as_uri()
                cell.font = Font(color='1261A0', underline='single')
        sheet.row_dimensions[index].height = min(360, max(72, ((len(proposal.get('risks', '')) // 35) + 1) * 16))
    for index, width in enumerate((14, 29, 25, 12, 18, 44, 44, 28, 35, 14, 45, 75), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    validation = DataValidation(type='list', formula1='"1,是"', allow_blank=True)
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f'A5:A{max(5, sheet.max_row)}')
    sheet.freeze_panes = 'F5'
    sheet.column_dimensions['B'].hidden = True  # 候选身份供程序校验，核对时无需阅读。
    sheet.auto_filter.ref = f'A4:L{max(4, sheet.max_row)}'
    sheet.sheet_view.showGridLines = False


def read_candidates(path):
    if path.suffix.lower() != '.xlsx':
        return []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if '候选配对' not in workbook.sheetnames:
            return []  # 旧版清单仍可继续使用。
        sheet = workbook['候选配对']
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=4, max_row=4))]
        if headers != [label for _, label in CANDIDATE_COLUMNS]:
            raise ValueError('候选配对列已改变，请重新生成核对表')
        result = []
        for cells in sheet.iter_rows(min_row=5, max_col=len(CANDIDATE_COLUMNS)):
            if not any(cell.value is not None for cell in cells):
                continue
            if any(cell.data_type == 'f' for cell in cells):
                raise ValueError('候选配对不支持公式，请填写文本')
            result.append({key: '' if cell.value is None else str(cell.value)
                           for (key, _), cell in zip(CANDIDATE_COLUMNS, cells)})
        return result
    finally:
        workbook.close()


def selected_candidates(path, trusted):
    by_id = {p['id']: p for p in trusted}
    selected, claimed = [], set()
    for row in read_candidates(path):
        if row['confirmed'].strip().lower() not in YES:
            continue
        proposal = by_id.get(row['id'])
        if proposal is None:
            raise ValueError('候选已失效或编号被修改，请重新识别刷新候选：' + row['id'])
        # 展示字段也不能被改造成另一笔材料；选择必须对应当前扫描证据。
        if any(row[key] != str(proposal[key]) for key, _ in CANDIDATE_COLUMNS if key != 'confirmed'):
            raise ValueError('候选内容已修改；请在核对清单修改字段并确认：' + row['id'])
        sources = set(proposal['members'])
        if claimed & sources:
            raise ValueError('多个已确认候选争用同一份原件；每张发票和订单只能选择一组')
        claimed.update(sources)
        selected.append(proposal)
    return selected
