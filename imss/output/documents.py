"""Template fields stay editable; supporting materials are inserted as full-page images."""
import io
from pathlib import Path
import re
from copy import deepcopy
from PIL import Image

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor

from imss.recognition.extract import material_pages
from imss.core.models import Bundle, MATERIALS, STANDALONE_SHEETS, Settings

MARKERS = {"invoice_images": "电子发票", "purchase_images": "购买记录", "promise_images": "承诺书", "all_materials": "all"}


def validate_signature(path: Path):
    if not path.is_file():
        raise ValueError(f"签字图片不存在: {path}")
    with Image.open(path) as image:
        if image.format not in ('PNG', 'JPEG', 'BMP', 'GIF', 'TIFF'):
            raise ValueError("签字图片请使用 PNG、JPG、BMP、GIF 或 TIFF")
        image.verify()


def replace_signature(paragraph, path: Path | None):
    if not re.search(r'{{\s*signature_image\s*}}', paragraph.text):
        return
    marker = '\ufffc'
    replace_fields(paragraph, {'signature_image': marker if path else ''})
    if not path:
        return
    with Image.open(path) as image:
        width, height = image.size
    factor = min(23 / width, 7 / height)
    for run in list(paragraph.runs):
        if marker not in run.text:
            continue
        parts = run.text.split(marker)
        run.text = parts[0]
        previous = run
        for following in parts[1:]:
            picture = paragraph.add_run()
            picture.font.size = Pt(1)
            previous._r.addnext(picture._r)
            picture.add_picture(str(path), width=Mm(width * factor), height=Mm(height * factor))
            tail = paragraph.add_run(following)
            if run._r.rPr is not None:
                tail._r.insert(0, deepcopy(run._r.rPr))
            picture._r.addnext(tail._r)
            previous = tail


def paragraphs_in(container):
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                yield from paragraphs_in(cell)


def replace_fields(paragraph, fields):
    # Replace across split Word runs without erasing the paragraph or its drawings.
    pattern = re.compile(r"{{\s*(" + "|".join(re.escape(key) for key in fields) + r")\s*}}")
    for match in reversed(list(pattern.finditer(paragraph.text))):
        position = 0
        started = False
        for run in paragraph.runs:
            text = run.text
            start, end = position, position + len(text)
            position = end
            if end <= match.start() or start >= match.end():
                continue
            left, right = max(0, match.start() - start), min(len(text), match.end() - start)
            replacement = str(fields[match.group(1)]) if not started else ""
            run.text = text[:left] + replacement + text[right:]
            started = True


def create_template(path: Path):
    if path.exists():
        raise FileExistsError(f"拒绝覆盖模板: {path}")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(18)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "宋体"
        style.font.color.rgb = RGBColor(0, 0, 0)
        from docx.oxml.ns import qn
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    document.styles["Normal"].font.size = Pt(11)
    document.add_paragraph("报销材料汇编", "Title")
    document.add_paragraph("序号：{{sequence}}    物品：{{item_name}}    金额：{{amount}} 元")
    document.add_paragraph("本文件汇集该笔支出的发票、购买记录及承诺书（如有），供报销核对。")
    document.add_paragraph("{{invoice_images}}")
    document.add_paragraph("{{purchase_images}}")
    document.add_paragraph("{{promise_images}}")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def build_document(bundle: Bundle, sequence: str, template: Path | None, destination: Path, settings: Settings, details=None):
    """按报销类别选模板，填可编辑表格，再渲染原始附件；签领表单独导出。

    details 的逐笔用途和人员信息已由流程层合并。签名只填指定占位符，
    模板承诺原图不改动；图片区布局与通用多页模板分开处理。
    """
    if destination.exists():
        raise FileExistsError(f"拒绝覆盖: {destination}")
    details = details or {}
    bundle = Bundle(bundle.key, bundle.item, bundle.amount,
                    [m for m in bundle.materials if m.kind not in STANDALONE_SHEETS], bundle.issues)
    signature = Path(details['signature_image']) if details.get('signature_image') else None
    if signature:
        validate_signature(signature)
    if not template:
        category = details.get("category") or next((m.category for m in bundle.materials if m.kind == "电子发票"), "物品") or "物品"
        derived = Path(__file__).resolve().parents[1] / "resources" / "templates" / (category + ".docx")
        if derived.is_file():
            template = derived
    document = Document(template) if template else Document()
    if not template:
        section = document.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(18)
        section.top_margin = section.bottom_margin = Mm(18)
        document.add_paragraph(f"{sequence} {bundle.item}", "Title")
        document.add_paragraph(f"报销金额：{bundle.amount} 元")
        document.add_paragraph("{{all_materials}}")
    invoice = next((m for m in bundle.materials if m.kind == "电子发票"), None)
    prize = details.get("is_prize")
    equal = details.get("amount_equal")
    fields = {"sequence": sequence, "item_name": bundle.item, "amount": bundle.amount,
              "invoice_number": details.get("invoice_number", invoice.invoice_ids[0] if invoice and invoice.invoice_ids else ""),
              "purpose": details.get("purpose", bundle.item), "student_id": details.get("student_id", ""),
              "passenger_name": details.get("passenger_name", ""),
              "passenger_name_2": details.get("passenger_name_2", ""),
              "passenger_name_3": details.get("passenger_name_3", ""),
              "phone": details.get("phone", ""), "actual_amount": details.get("actual_amount", ""),
              "prize_yes": "☑" if prize is True else "□", "prize_no": "☑" if prize is False else "□",
              "equal_yes": "☑" if equal is True else "□", "equal_no": "☑" if equal is False else "□"}
    containers = [document]
    for section in document.sections:
        containers.extend([section.header, section.footer, section.first_page_header, section.first_page_footer,
                           section.even_page_header, section.even_page_footer])
    for container in containers:
        for paragraph in paragraphs_in(container):
            replace_fields(paragraph, fields)
            replace_signature(paragraph, signature)
    panels = [(p, p.text.strip()) for p in paragraphs_in(document) if p.text.strip() in ("{{invoice_panel}}", "{{evidence_panel}}")]
    if panels:
        if sorted(text for _, text in panels) != ['{{evidence_panel}}', '{{invoice_panel}}']:
            raise ValueError("单页模板需要唯一的 invoice_panel 和 evidence_panel")
        for marker, panel in panels:
            selected = [m for m in bundle.materials if (m.kind == "电子发票") == (panel == "{{invoice_panel}}")]
            insert_panel(document, marker, selected, settings)
        destination.parent.mkdir(parents=True, exist_ok=True)
        document.save(destination)
        return
    markers = []
    for paragraph in paragraphs_in(document):
        match = re.fullmatch(r"\s*{{\s*(\w+)\s*}}\s*", paragraph.text)
        if match and match.group(1) in MARKERS:
            markers.append((paragraph, MARKERS[match.group(1)]))
        elif any(re.search(r"{{\s*" + key + r"\s*}}", paragraph.text) for key in MARKERS):
            raise ValueError("图片占位符必须独占一段，不能与文字混排")
    marker_kinds = [kind for _, kind in markers]
    if len(marker_kinds) != len(set(marker_kinds)) or ("all" in marker_kinds and len(marker_kinds) > 1):
        raise ValueError("模板图片占位符重复，或 all_materials 与分类占位符混用")
    # Existing templates with no image markers still receive all attachments at the end.
    if not markers:
        markers = [(document.add_paragraph(), "all")]
    used = set()
    for marker, kind in markers:
        selected = [m for m in bundle.materials if kind == "all" or m.kind == kind]
        selected.sort(key=lambda m: (MATERIALS.index(m.kind), m.source))
        for material in selected:
            insert_material(document, marker, material, settings)
            used.add(material.source)
        marker._element.getparent().remove(marker._element)
    # A template may contain only one of the category markers. Never drop other materials.
    for material in sorted(bundle.materials, key=lambda m: (MATERIALS.index(m.kind), m.source)):
        if material.source not in used:
            marker = document.add_paragraph()
            insert_material(document, marker, material, settings)
            marker._element.getparent().remove(marker._element)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(".partial.docx")
    try:
        document.save(temporary)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)


def insert_material(document, marker, material, settings):
    # Use the smallest printable area in the template so mixed sections cannot overflow.
    width = min(s.page_width - s.left_margin - s.right_margin for s in document.sections)
    height = min(s.page_height - s.top_margin - s.bottom_margin for s in document.sections) - Mm(20)
    cell = marker._element.getparent()
    if cell.tag.endswith("}tc"):
        from docx.oxml.ns import qn
        cell_width = cell.find("./" + qn("w:tcPr") + "/" + qn("w:tcW"))
        if cell_width is not None and cell_width.get(qn("w:type")) == "dxa":
            width = min(width, int(cell_width.get(qn("w:w"))) * 635 - Mm(4))
    if width <= 0 or height <= 0:
        raise ValueError("模板可用页面尺寸过小")
    count = 0
    for index, image in enumerate(material_pages(Path(material.archived), material.extension, settings)):
        try:
            caption = document.add_paragraph()
            caption.paragraph_format.page_break_before = True
            caption.paragraph_format.keep_with_next = True
            caption.paragraph_format.space_after = Pt(4)
            caption.add_run(f"{material.kind}  第 {index + 1} 页").bold = True
            marker._element.addprevious(caption._element)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            factor = min(width / image.width, height / image.height)
            buffer = io.BytesIO()
            image.save(buffer, format="PNG")
            buffer.seek(0)
            paragraph.add_run().add_picture(buffer, width=int(image.width * factor), height=int(image.height * factor))
            marker._element.addprevious(paragraph._element)
            count += 1
        finally:
            image.close()
    if not count:
        raise ValueError(f"材料没有可插入页面: {material.source}")


def insert_panel(document, marker, materials, settings):
    """Compact invoice/evidence rows reproduce the observed one-page form family."""
    pages = []
    try:
        for material in sorted(materials, key=lambda m: (MATERIALS.index(m.kind), m.source)):
            for image in material_pages(Path(material.archived), material.extension, settings):
                pages.append(image)
        if not pages:
            marker._element.getparent().remove(marker._element)
            return
        if len(pages) > 6:
            raise ValueError("单页模板每个图片区最多容纳 6 页；请拆分该组或使用多页模板")
        section = document.sections[0]
        width = section.page_width - section.left_margin - section.right_margin
        table = document.add_table(rows=1, cols=len(pages))
        table.autofit = False
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        properties = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement('w:' + edge)
            border.set(qn('w:val'), 'nil')
            borders.append(border)
        properties.append(borders)
        for column in table.columns:
            column.width = int(width / len(pages))
        for cell, image in zip(table.rows[0].cells, pages):
            cell.width = int(width / len(pages))
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.keep_with_next = False
            factor = min((width / len(pages) - Mm(4)) / image.width, Mm(108) / image.height)
            stream = io.BytesIO()
            image.save(stream, format='PNG')
            stream.seek(0)
            paragraph.add_run().add_picture(stream, width=int(image.width * factor), height=int(image.height * factor))
        # Word merges adjacent tables without an intervening paragraph; that can
        # erase neighboring form borders and make pictures touch the declaration.
        for element in ("before", "table", "after"):
            if element == "table":
                marker._element.addprevious(table._tbl)
            else:
                spacer = document.add_paragraph()
                spacer.paragraph_format.space_before = spacer.paragraph_format.space_after = Pt(0)
                spacer.paragraph_format.line_spacing = Pt(2)
                spacer.add_run().font.size = Pt(1)
                marker._element.addprevious(spacer._element)
        marker._element.getparent().remove(marker._element)
    finally:
        for page in pages:
            page.close()
