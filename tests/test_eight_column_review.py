"""八列核对工作流：覆盖/回退、状态筛选、重配对、原件链接及升级。"""
from copy import deepcopy
from contextlib import redirect_stdout
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image

from imss.core.models import Material, Settings
from imss.output.review import write_review, read_review, apply_review
from imss.output.review_table import COLUMNS
from imss.workflow import scan, build, update_review, recheck
from imss.storage import digest_file
from main import main


class EightColumnReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.source = self.root / '原始 输入'
        self.source.mkdir()

    def scanned(self, texts=None):
        texts = texts or {'i.png': '电子发票\n物品名称：打印纸\n小写￥35.80',
                          'p.png': '交易成功\n商品名称：打印纸\n实付款￥35.80'}
        for index, name in enumerate(texts):
            path = self.source / name
            path.parent.mkdir(parents=True, exist_ok=True)
            Image.new('RGB', (120 + index, 100), 'white').save(path)
        with patch('imss.workflow.extract', side_effect=lambda path, *_: (texts[path.relative_to(self.source).as_posix()], [])):
            manifest = scan(self.source, None, Settings(workers=1))
        return manifest, manifest.with_suffix('.xlsx')

    def edit(self, path, edits):
        book = load_workbook(path)
        for cells in book.active.iter_rows(min_row=5):
            updates = edits.get(cells[0].value, {})
            for column, value in updates.items():
                cells[column - 1].value = value
        book.save(path)
        book.close()

    def test_eight_columns_one_sheet_blank_overrides_and_defaults(self):
        manifest, path = self.scanned()
        book = load_workbook(path)
        self.assertEqual(book.sheetnames, ['核对清单'])
        sheet = book.active
        self.assertEqual(sheet.max_column, 8)
        self.assertEqual([c.value for c in sheet[4]], [label for _, label in COLUMNS])
        for cells in sheet.iter_rows(min_row=5):
            self.assertEqual([c.value for c in cells[4:7]], [None, None, None])
            self.assertEqual(cells[7].value, '已校验')
            self.assertIn('识别金额：35.80', cells[2].value)
            self.assertIn('打印纸', cells[2].value)
            self.assertEqual(cells[0].fill.fgColor.rgb[-6:], 'E2F0D9')
        self.assertTrue(any(v.formula1 == '"已校验,待校验"' and str(v.sqref) == 'H5:H6'
                            and v.showDropDown is False for v in sheet.data_validations.dataValidation))
        book.close()
        self.assertFalse(build(manifest, skip_pdf=True)['pending'])

    def test_problem_rows_default_waiting(self):
        manifest, path = self.scanned({'i.png': '电子发票\n物品名称：打印纸\n小写￥35.80',
                                       'p.png': '交易成功\n实付款￥35.80'})
        self.assertTrue(all(row['confirmed'] == '待校验' for row in read_review(path)))
        report = build(manifest, skip_pdf=True)
        self.assertFalse(report['generated'])
        self.assertEqual(len(report['pending']), 2)

    def test_blank_manual_fields_keep_recognition_and_clearing_reverts(self):
        manifest, path = self.scanned()
        edits = {name: {6: '活动用纸', 7: '36.00'} for name in ('i.png', 'p.png')}
        self.edit(path, edits)
        report = build(manifest, skip_pdf=True)
        self.assertEqual(report['generated'][0]['item'], '活动用纸')
        self.assertEqual(report['generated'][0]['amount'], '36.00')
        self.edit(path, {name: {6: None, 7: None} for name in edits})
        report = build(manifest, skip_pdf=True)
        self.assertEqual(report['generated'][0]['item'], '打印纸')
        self.assertEqual(report['generated'][0]['amount'], '35.80')
        self.assertFalse((Path(report['output']) / '01.活动用纸_36.00.docx').exists())
        data = json.loads(manifest.read_text(encoding='utf-8'))
        self.assertTrue(all(m['item'] == '打印纸' and m['amount'] == '35.80' for m in data['materials']))

    def test_cli_excel_entry_uses_latest_excel_changes_without_editing_scan_json(self):
        manifest, path = self.scanned({'i.png': '电子发票\n物品名称：打印纸\n小写￥35.80',
                                       'p.png': '模糊的商品截图'})
        scanned = manifest.read_bytes()
        self.edit(path, {'i.png': {6: '活动用纸', 7: '36', 8: '已校验'},
                         'p.png': {5: '购买记录', 6: '活动用纸', 7: '36', 8: '已校验'}})
        self.assertEqual(manifest.read_bytes(), scanned)
        with patch('imss.workflow.extract') as extract, redirect_stdout(io.StringIO()):
            self.assertEqual(main(['build', str(path), '--skip-pdf']), 0)
        extract.assert_not_called()
        data = json.loads(manifest.read_text(encoding='utf-8'))
        self.assertEqual({m['source']: m['kind'] for m in data['materials']},
                         {'i.png': '电子发票', 'p.png': '未知'})
        report = data['report']
        self.assertEqual((report['generated'][0]['item'], report['generated'][0]['amount']),
                         ('活动用纸', '36.00'))
        self.assertEqual(Path(report['review_file']), path.resolve())
        self.assertTrue((Path(report['output']) / '01.活动用纸_36.00.docx').is_file())
        self.edit(path, {name: {6: '办公用纸', 7: '37'} for name in ('i.png', 'p.png')})
        with redirect_stdout(io.StringIO()):
            self.assertEqual(main(['build', str(path), '--skip-pdf']), 0)
        output = Path(report['output'])
        self.assertTrue((output / '01.办公用纸_37.00.docx').is_file())
        self.assertFalse((output / '01.活动用纸_36.00.docx').exists())

    def test_excel_entry_reads_the_selected_workbook_instead_of_default_workbook(self):
        manifest, path = self.scanned()
        original = path.read_bytes()
        selected = path.with_name('另存的人工核对.xlsx')
        selected.write_bytes(original)
        self.edit(selected, {name: {6: '人工修改名称', 7: '36'} for name in ('i.png', 'p.png')})
        report = build(selected, skip_pdf=True)
        self.assertEqual(report['generated'][0]['item'], '人工修改名称')
        self.assertEqual(report['generated'][0]['amount'], '36.00')
        self.assertEqual(path.read_bytes(), original)
        self.assertTrue(manifest.is_file())

    def test_excel_entry_rejects_conflicting_review_without_changing_files(self):
        manifest, path = self.scanned()
        before = (manifest.read_bytes(), path.read_bytes())
        with self.assertRaisesRegex(ValueError, '--review 不一致'):
            build(path, review=path.with_name('other.xlsx'), skip_pdf=True)
        self.assertEqual((manifest.read_bytes(), path.read_bytes()), before)

    def test_missing_automatic_scan_record_reports_how_to_restore_it(self):
        manifest, path = self.scanned()
        original = path.read_bytes()
        manifest.unlink()
        with self.assertRaisesRegex(ValueError, '缺少程序保存的扫描记录'):
            build(path, skip_pdf=True)
        self.assertEqual(path.read_bytes(), original)

    def test_workbook_and_pending_cli_instructions_use_excel_entry(self):
        _, path = self.scanned()
        book = load_workbook(path)
        self.assertIn(f'build "{path.resolve()}"', book.active['A3'].value)
        self.assertNotIn('.json', book.active['A3'].value)
        book.close()
        self.edit(path, {name: {8: '待校验'} for name in ('i.png', 'p.png')})
        output = io.StringIO()
        with redirect_stdout(output):
            self.assertEqual(main(['build', str(path), '--skip-pdf']), 2)
        self.assertIn(f'build "{path}"', output.getvalue())
        self.assertNotIn('build "' + str(path.with_suffix('.json')), output.getvalue())

    def test_scan_cli_instructions_use_excel_entry(self):
        path = self.source / 'sample.png'
        Image.new('RGB', (120, 100), 'white').save(path)
        output = io.StringIO()
        with patch('imss.workflow.extract', return_value=('一份待核对材料', [])), redirect_stdout(output):
            self.assertEqual(main(['scan', str(self.source)]), 0)
        review = self.root / '核对清单' / '核对清单.xlsx'
        self.assertIn(f'build "{review}"', output.getvalue())
        self.assertNotIn('build "' + str(review.with_suffix('.json')), output.getvalue())

    def test_verified_selection_is_recomputed_and_waiting_rows_excluded(self):
        manifest, path = self.scanned()
        self.assertEqual(len(build(manifest, skip_pdf=True)['generated']), 1)
        self.edit(path, {'p.png': {8: '待校验'}})
        report = build(manifest, skip_pdf=True)
        self.assertFalse(report['generated'])
        self.assertTrue(any('待校验' in str(entry) for entry in report['pending']))
        self.assertFalse(any(Path(report['output']).rglob('*.docx')))
        self.edit(path, {'p.png': {8: '已校验'}})
        self.assertEqual(len(build(manifest, skip_pdf=True)['generated']), 1)

    def test_repair_missing_item_then_verify_and_generate(self):
        manifest, path = self.scanned({'i.png': '电子发票\n物品名称：打印纸\n小写￥35.80',
                                       'p.png': '交易成功\n实付款￥35.80'})
        self.edit(path, {'i.png': {8: '已校验'}, 'p.png': {6: '打印纸', 8: '已校验'}})
        self.assertFalse(build(manifest, skip_pdf=True)['pending'])

    def test_changed_names_repair_sources_instead_of_reusing_old_pairs(self):
        texts = {'a-i.png': '电子发票\n物品名称：纸杯\n小写￥35.80',
                 'b-i.png': '电子发票\n物品名称：打印纸\n小写￥35.80',
                 'a-p.png': '交易成功\n商品名称：纸杯\n实付款￥35.80',
                 'b-p.png': '交易成功\n商品名称：打印纸\n实付款￥35.80'}
        manifest, path = self.scanned(texts)
        self.assertEqual(len(build(manifest, skip_pdf=True)['generated']), 2)
        self.edit(path, {'a-p.png': {6: '打印纸'}, 'b-p.png': {6: '纸杯'}})
        report = build(manifest, skip_pdf=True)
        self.assertFalse(report['pending'])
        self.assertEqual({r['item']: set(r['sources']) for r in report['generated']},
                         {'纸杯': {'a-i.png', 'b-p.png'}, '打印纸': {'b-i.png', 'a-p.png'}})

    def test_verified_is_not_permission_to_bypass_conflicting_totals(self):
        manifest, path = self.scanned()
        self.edit(path, {'p.png': {7: 99}})
        report = build(manifest, skip_pdf=True)
        self.assertFalse(report['generated'])
        self.assertTrue(report['pending'])
        self.assertTrue(all(row['confirmed'] == '已校验' for row in read_review(path)))

    def test_unique_verified_split_payments_can_be_repaired_without_group_column(self):
        manifest, path = self.scanned({'i.png': '电子发票\n物品名称：活动用品\n小写￥480',
                                       'p1.png': '交易成功\n商品名称：活动用品\n实付款￥80',
                                       'p2.png': '交易成功\n商品名称：活动用品\n实付款￥400'})
        self.edit(path, {name: {8: '已校验'} for name in ('i.png', 'p1.png', 'p2.png')})
        report = build(manifest, skip_pdf=True)
        self.assertFalse(report['pending'])
        self.assertEqual(len(report['generated']), 1)
        self.assertEqual(len(report['generated'][0]['sources']), 3)

    def test_unicode_space_hash_and_percent_file_links_target_existing_originals(self):
        manifest, path = self.scanned({'子 目录/发票 #1 50%.png': '电子发票\n物品名称：打印纸\n小写￥35.80',
                                       '购买记录 & +.png': '交易成功\n商品名称：打印纸\n实付款￥35.80'})
        book = load_workbook(path)
        for cells in book.active.iter_rows(min_row=5):
            target = Path(cells[0].hyperlink.target)
            self.assertEqual(target, (self.source / cells[0].value).resolve())
            self.assertTrue(target.is_file())
            self.assertNotIn('file:///', cells[0].hyperlink.target)
            self.assertNotIn('%E', cells[0].hyperlink.target)
        book.close()
        with ZipFile(path) as archive:
            links = archive.read('xl/worksheets/_rels/sheet1.xml.rels').decode('utf-8')
            self.assertIn('原始 输入', links)
            self.assertIn('50%', links)
        self.assertFalse(build(manifest, skip_pdf=True)['pending'])

    def test_formulas_duplicate_or_deleted_source_rows_and_invalid_status_are_rejected(self):
        for change, message in [({6: '=1+1'}, '公式'), ({8: '1'}, '校验标志位'),
                                ({7: '1.234'}, '两位小数'), ({5: '不存在的类型'}, '类型无效')]:
            with self.subTest(change=change):
                manifest, path = self.scanned()
                self.edit(path, {'i.png': change})
                with self.assertRaisesRegex(ValueError, message):
                    build(manifest, skip_pdf=True)
        manifest, path = self.scanned()
        book = load_workbook(path)
        book.active.delete_rows(5)
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, '不可删除行'):
            build(manifest, skip_pdf=True)

    def test_refresh_preserves_pending_edits_and_original_ocr(self):
        manifest, path = self.scanned()
        self.edit(path, {'p.png': {6: '待核查名称', 7: '尚未确认的金额', 8: '待校验'}})
        build(manifest, skip_pdf=True)
        row = next(r for r in read_review(path) if r['source'] == 'p.png')
        self.assertEqual((row['item'], row['amount'], row['confirmed']), ('待核查名称', '尚未确认的金额', '待校验'))
        self.assertIn('识别名称：打印纸', row['ocr'])
        self.assertNotIn('待核查名称', row['ocr'])

    def test_review_upgrade_does_not_reocr_or_change_json_and_can_be_repeated(self):
        manifest, path = self.scanned()
        before = manifest.read_bytes()
        self.edit(path, {'p.png': {6: '人工名称', 8: '待校验'}})
        with patch('imss.workflow.extract') as extract:
            update_review(manifest)
            update_review(manifest)
        extract.assert_not_called()
        self.assertEqual(before, manifest.read_bytes())
        self.assertEqual(next(r for r in read_review(path) if r['source'] == 'p.png')['item'], '人工名称')

    def test_legacy_migration_keeps_edits_status_and_rewrites_broken_links(self):
        manifest, path = self.scanned()
        from tests.legacy_review_fixture import write_review as old_write
        materials = [Material(**m) for m in json.loads(manifest.read_text(encoding='utf-8'))['materials']]
        old_write(path, materials, self.source)
        book = load_workbook(path)
        for cells in book.active.iter_rows(min_row=5):
            cells[2].value = '1'
            cells[3].value = '活动用纸'
            cells[7].hyperlink = 'file:///D:/wrong/%E4%B8%AD.png'
        book.save(path)
        book.close()
        update_review(manifest)
        self.assertTrue(all(r['item'] == '活动用纸' and r['confirmed'] == '已校验' for r in read_review(path)))
        self.assertFalse(build(manifest, skip_pdf=True)['pending'])
        book = load_workbook(path)
        self.assertEqual(book.sheetnames, ['核对清单'])
        self.assertEqual(book.active.max_column, 8)
        self.assertTrue(Path(book.active['A5'].hyperlink.target).is_file())
        book.close()

    def test_legacy_numeric_format_does_not_become_a_user_override_or_unverify_good_rows(self):
        manifest, path = self.scanned()
        from tests.legacy_review_fixture import write_review as old_write
        materials = [Material(**m) for m in json.loads(manifest.read_text(encoding='utf-8'))['materials']]
        old_write(path, materials, self.source)
        update_review(manifest)
        rows = read_review(path)
        self.assertTrue(all(row['confirmed'] == '已校验' for row in rows))
        self.assertTrue(all(row['amount'] == '' for row in rows))

    def test_recheck_keeps_user_edits_but_clearing_returns_to_recognized_value(self):
        manifest, path = self.scanned()
        self.edit(path, {name: {6: '活动用纸'} for name in ('i.png', 'p.png')})
        with patch('imss.workflow.extract') as extract:
            recheck(manifest)
        extract.assert_not_called()
        self.edit(path, {name: {6: None} for name in ('i.png', 'p.png')})
        report = build(manifest, skip_pdf=True)
        self.assertEqual(report['generated'][0]['item'], '打印纸')

    def test_manual_type_and_name_override_renames_standalone_sheet(self):
        manifest, path = self.scanned({'unknown.png': '一份尚未识别的表单'})
        self.edit(path, {'unknown.png': {5: '签到表', 6: '迎新交流活动签到表', 8: '已校验'}})
        report = build(manifest, skip_pdf=True)
        self.assertFalse(report['pending'])
        self.assertTrue((Path(report['output']) / '01.迎新交流活动签到表.pdf').is_file())

    def test_ocr_non_xml_characters_are_filtered_without_changing_source_text(self):
        material = Material('i.png', '', '.png', kind='未知', text='识别\ufffe\ud800结果\x00')
        path = self.root / 'review.xlsx'
        write_review(path, [material])
        self.assertIn('识别结果', read_review(path)[0]['ocr'])
        self.assertIn('控制字符', read_review(path)[0]['ocr'])
        self.assertIn('\ufffe', material.text)
