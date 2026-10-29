from pathlib import Path
import os
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image, ImageDraw, ImageFont
from docx import Document
import pypdfium2 as pdfium

from imss.recognition.extract import ocr_image, tesseract_path
from imss.core.models import Settings
from imss.runtime.office import convert_office, run_process, soffice_path
from imss.runtime.paths import tessdata_path


class LocalIntegrationTests(unittest.TestCase):
    @patch.dict(os.environ, {'PATH': '', 'TESSDATA_PREFIX': 'nonexistent-system-models'})
    def test_real_ocr(self):
        settings = Settings(languages='eng')
        try:
            executable = tesseract_path(settings)
            args = [executable, '--list-langs']
            data = tessdata_path(settings)
            if data:
                args.extend(['--tessdata-dir', str(data)])
            result = run_process(args, 20)
            if 'eng' not in result.stdout.decode('utf-8', errors='replace').splitlines()[1:]:
                self.skipTest('本机缺少英文 OCR 模型')
        except RuntimeError as exc:
            self.skipTest(str(exc))
        image = Image.new('RGB', (600, 180), 'white')
        draw = ImageDraw.Draw(image)
        font = None
        for candidate in (Path('C:/Windows/Fonts/arial.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')):
            if candidate.is_file():
                font = ImageFont.truetype(str(candidate), 50)
                break
        if font is None:
            self.skipTest('缺少用于 OCR 集成测试的 Arial/DejaVuSans 字体')
        draw.text((25, 55), 'TOTAL 35.80', font=font, fill='black')
        text, issues = ocr_image(image, settings)
        self.assertIn('35.80', text)
        self.assertFalse(issues)

    @patch.dict(os.environ, {'PATH': '', 'TESSDATA_PREFIX': 'nonexistent-system-models'})
    def test_real_libreoffice_export(self):
        settings = Settings(office_backend='libreoffice')
        try:
            soffice_path(settings)
        except RuntimeError as exc:
            self.skipTest(str(exc))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            document = Document()
            document.add_paragraph('Offline export verification 35.80')
            source = root / 'check.docx'
            document.save(source)
            target = convert_office(source, root / 'check.pdf', settings)
            with pdfium.PdfDocument(target) as pdf:
                self.assertEqual(len(pdf), 1)


if __name__ == '__main__':
    unittest.main()
