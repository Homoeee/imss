"""核对闭环回归：建议不能直接报销，人工修改必须真正影响最终材料。"""
from copy import deepcopy
from itertools import permutations
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from PIL import Image
from docx import Document

from main import main
from imss.core.matching import group_materials
from imss.core.models import Material, Settings
from imss.output.review import apply_review, read_review, review_rows, write_review
from imss.output.workspace import default_output
from imss.storage import digest_file
from imss.workflow import build, scan


def material(source, kind='电子发票', item='打印纸', amount='35.80', **kwargs):
    return Material(source, source, '.png', kind=kind, item=item, amount=amount, **kwargs)


class AmountFallbackTests(unittest.TestCase):
    def test_unique_amount_suggests_name_but_blocks_generation(self):
        invoice = material('i')
        order = material('p', '购买记录', item='', issues=['未找到明确的物品名称'])
        bundles, pending = group_materials([invoice, order])
        self.assertFalse(pending)
        self.assertEqual(order.item, '打印纸')
        self.assertEqual(order.match_basis, 'amount_only')
        self.assertTrue(bundles[0].issues)
        self.assertTrue(all(row['severity'] == '待确认' for row in review_rows([invoice, order])))

    def test_equal_price_invoices_are_not_guessed(self):
        order = material('p', '购买记录', item='')
        bundles, pending = group_materials([material('i1'), material('i2', item='水杯'), order])
        self.assertEqual(pending, [order])
        self.assertEqual(order.item, '')
        self.assertTrue(all(order not in b.materials for b in bundles))

    def test_equal_price_orders_cannot_claim_the_same_invoice(self):
        for sequence in permutations([material('i'), material('p1', '购买记录', item=''), material('p2', '购买记录', item='')]):
            bundles, pending = group_materials(deepcopy(sequence))
            self.assertEqual({m.source for m in pending}, {'p1', 'p2'})
            self.assertEqual(len(bundles[0].materials), 1)

    def test_strong_pairing_finishes_before_amount_fallback(self):
        records = [material('i1'), material('i2', item='水杯'),
                   material('a-unknown', '购买记录', item=''), material('z-known', '购买记录')]
        for sequence in (records, list(reversed(records))):
            bundles, pending = group_materials(deepcopy(sequence))
            self.assertFalse(pending)
            cup = next(b for b in bundles if b.item == '水杯')
            self.assertEqual({m.source for m in cup.materials}, {'i2', 'a-unknown'})
            self.assertTrue(cup.issues)

    def test_occupied_invoice_and_conflicting_ids_are_excluded(self):
        order = material('p2', '购买记录', item='')
        _, pending = group_materials([material('i'), material('p1', '购买记录'), order])
        self.assertIn(order, pending)
        for field in ('order_ids', 'invoice_ids'):
            with self.subTest(field=field):
                order = material('p', '购买记录', item='', **{field: ['222222']})
                _, pending = group_materials([material('i', **{field: ['111111']}), order])
                self.assertIn(order, pending)

    def test_nonblank_conflicting_item_never_uses_amount_fallback(self):
        order = material('p', '购买记录', item='水杯')
        _, pending = group_materials([material('i'), order])
        self.assertEqual(pending, [order])
        self.assertEqual(order.item, '水杯')


class ReviewWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / 'input'
        self.source.mkdir()

    def tearDown(self):
        self.temp.cleanup()

    def make_scan(self, amount_only=True):
        texts = {'i.png': '电子发票\n物品名称：打印纸\n（小写）￥35.80',
                 'p.png': '交易成功\n实付款：35.80' if amount_only else '交易成功\n物品名称：打印纸\n实付款：35.80'}
        for index, name in enumerate(texts):
            Image.new('RGB', (120 + index, 100), 'white').save(self.source / name)
        with patch('imss.workflow.extract', side_effect=lambda path, *_: (texts[path.name], [])):
            manifest = scan(self.source, None, Settings(workers=1))
        return manifest, manifest.with_suffix('.xlsx')

    def edit_rows(self, path, edit):
        workbook = load_workbook(path)
        for row in workbook.active.iter_rows(min_row=5):
            edit(row)
        workbook.save(path)
        workbook.close()

    def test_new_extraction_conflicts_are_red_even_in_known_group(self):
        for issue in ('实付金额冲突: 12.70, 2.30', '包含多个商品或候选标题: 纸、笔',
                      '金额字段无法完整读取，请核对原件', '订单尚未付款或已经关闭，不能作为已付款凭证'):
            with self.subTest(issue=issue):
                records = [material('i', group='01'),
                           material('p', '购买记录', group='01', issues=[issue])]
                self.assertTrue(all(row['severity'] == '严重' for row in review_rows(records)))

    def test_review_confirmation_unlocks_generation_and_rebuild_uses_latest_edits(self):
        manifest, review = self.make_scan()
        original = {p.name: digest_file(p) for p in self.source.iterdir()}
        first = build(manifest, skip_pdf=True)
        self.assertTrue(first['pending'])
        self.assertFalse(list(default_output(self.source).iterdir()))
        self.edit_rows(review, lambda row: setattr(row[7], 'value', '已校验'))
        confirmed = build(manifest, skip_pdf=True)
        self.assertFalse(confirmed['pending'])
        self.assertTrue(all(row['severity'] == '正常' for row in read_review(review)))
        output = default_output(self.source)
        self.assertTrue((output / '01.打印纸_35.80.docx').is_file())
        self.edit_rows(review, lambda row: setattr(row[5], 'value', '活动用纸'))
        renamed = build(manifest, skip_pdf=True)
        self.assertFalse(renamed['pending'])
        self.assertFalse((output / '01.打印纸_35.80.docx').exists())
        doc = Document(output / '01.活动用纸_35.80.docx')
        self.assertIn('活动用纸', '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells))
        self.assertEqual(original, {p.name: digest_file(p) for p in self.source.iterdir()})
        self.assertEqual({p.name for p in self.root.iterdir()}, {'input', '报销材料', '核对清单'})
        self.assertEqual({p.name for p in manifest.parent.iterdir()}, {'核对清单.json', '核对清单.xlsx'})
        self.assertEqual({p.name for p in output.iterdir() if p.is_dir()}, {'原始材料'})
        self.assertEqual({p.name for p in (output / '原始材料').iterdir()},
                         {'01.活动用纸-35.80元-电子发票.png', '01.活动用纸-35.80元-购买记录.png'})

    def test_confirmed_rows_do_not_bypass_amount_validation(self):
        manifest, review = self.make_scan()
        def edit(row):
            row[7].value = '已校验'
            if row[0].value == 'p.png':
                row[6].value = '36.00'
        self.edit_rows(review, edit)
        report = build(manifest, skip_pdf=True)
        self.assertTrue(any('金额' in issue for entry in report['pending'] for issue in entry['issues']))
        self.assertFalse(list(default_output(self.source).iterdir()))
        self.assertTrue(all(row['severity'] == '严重' for row in read_review(review)))

    def test_manual_pairing_of_ambiguous_orders_controls_generated_sources(self):
        texts = {'paper-i.png': '电子发票\n物品名称：打印纸\n（小写）￥35.80',
                 'cup-i.png': '电子发票\n物品名称：水杯\n（小写）￥35.80',
                 'paper-p.png': '交易成功\n实付款：35.80', 'cup-p.png': '交易成功\n实付款：35.80'}
        for index, name in enumerate(texts):
            Image.new('RGB', (120 + index, 100), 'white').save(self.source / name)
        with patch('imss.workflow.extract', side_effect=lambda path, *_: (texts[path.name], [])):
            manifest = scan(self.source, None, Settings(workers=1))
        review = manifest.with_suffix('.xlsx')
        self.assertTrue(all(row['severity'] == '严重' for row in read_review(review)))
        def edit(row):
            prefix = row[0].value.split('-')[0]
            row[7].value = '已校验'
            row[5].value = '打印纸' if prefix == 'paper' else '水杯'
        self.edit_rows(review, edit)
        report = build(manifest, skip_pdf=True)
        self.assertFalse(report['pending'])
        self.assertEqual({entry['item']: set(entry['sources']) for entry in report['generated']},
                         {'打印纸': {'paper-i.png', 'paper-p.png'}, '水杯': {'cup-i.png', 'cup-p.png'}})
        self.assertTrue(all(row['severity'] == '正常' for row in read_review(review)))

    def test_partial_confirmation_does_not_approve_other_rows(self):
        manifest, review = self.make_scan()
        def edit(row):
            if row[0].value == 'p.png':
                row[7].value = '已校验'
        self.edit_rows(review, edit)
        self.assertTrue(build(manifest, skip_pdf=True)['pending'])

    def test_manual_changes_to_generated_files_block_rebuild(self):
        manifest, _ = self.make_scan(amount_only=False)
        report = build(manifest, skip_pdf=True)
        word = default_output(self.source) / report['generated'][0]['word']
        word.write_bytes(b'manually changed')
        with self.assertRaisesRegex(ValueError, '人工修改'):
            build(manifest, skip_pdf=True)
        self.assertEqual(word.read_bytes(), b'manually changed')

    def test_unconfirmed_edits_are_visible_and_cannot_generate_stale_values(self):
        manifest, review = self.make_scan(amount_only=False)
        self.edit_rows(review, lambda row: setattr(row[5], 'value', '活动用纸'))
        self.edit_rows(review, lambda row: setattr(row[7], 'value', '待校验'))
        report = build(manifest, skip_pdf=True)
        self.assertTrue(report['pending'])
        self.assertFalse(list(default_output(self.source).iterdir()))
        rows = read_review(review)
        self.assertTrue(all(row['item'] == '活动用纸' for row in rows))
        self.assertTrue(all('待校验' in row['issues'] for row in rows))

    def test_user_files_are_not_overwritten(self):
        manifest, _ = self.make_scan(amount_only=False)
        output = default_output(self.source)
        existing = output / '01.打印纸_35.80.docx'
        existing.write_bytes(b'user file')
        with self.assertRaisesRegex(ValueError, '拒绝覆盖'):
            build(manifest, skip_pdf=True)
        self.assertEqual(list(output.iterdir()), [existing])
        self.assertEqual(existing.read_bytes(), b'user file')

    def test_colors_sorting_links_and_roundtrip(self):
        good = [material('good-i'), material('good-p', '购买记录')]
        orange = [material('orange-i', amount='20.00'), material('orange-p', '购买记录', item='', amount='20.00')]
        yellow = material('yellow', '新闻稿', issues=['OCR 平均置信度偏低，请核对原件'])
        bad = material('bad', '未知', issues=['未识别材料类型'])
        records = [*good, *orange, yellow, bad]
        for record in records:
            (self.source / record.source).touch()
        path = self.root / '核对清单.xlsx'
        write_review(path, records, self.source)
        workbook = load_workbook(path)
        sheet = workbook.active
        self.assertEqual([sheet.cell(i, 1).value for i in range(5, 12)], ['bad', 'orange-i', 'orange-p', 'yellow', 'good-i', 'good-p', None])
        self.assertEqual([sheet.cell(i, 1).fill.fgColor.rgb[-6:] for i in (5, 6, 8, 9)], ['FADBD8', 'FCE4C5', 'FFF2CC', 'E2F0D9'])
        self.assertTrue(Path(sheet['A5'].hyperlink.target).is_file())
        self.assertEqual(sheet.freeze_panes, 'B5')
        self.assertTrue(sheet.auto_filter.ref)
        workbook.close()
        self.assertEqual(len(read_review(path)), len(records))

    def test_xlsx_rejects_missing_rows_and_formulas(self):
        manifest, review = self.make_scan()
        records = [Material(**m) for m in json.loads(manifest.read_text(encoding='utf-8'))['materials']]
        workbook = load_workbook(review)
        workbook.active['D5'] = '=1+1'
        workbook.save(review)
        workbook.close()
        with self.assertRaisesRegex(ValueError, '公式'):
            apply_review(review, deepcopy(records))
        write_review(review, records)
        workbook = load_workbook(review)
        workbook.active.delete_rows(5)
        workbook.save(review)
        workbook.close()
        with self.assertRaisesRegex(ValueError, '不可删除行'):
            apply_review(review, deepcopy(records))

    def test_text_beginning_with_equal_is_stored_as_text(self):
        path = self.root / '核对清单.xlsx'
        write_review(path, [material('i', item='=测试')])
        self.assertIn('识别名称：=测试', read_review(path)[0]['ocr'])


if __name__ == '__main__':
    unittest.main()
