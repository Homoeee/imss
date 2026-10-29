"""Configured person fields must not overwrite the template's promise image."""
import csv
from contextlib import redirect_stdout
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image
from docx import Document
from imss.output.documents import build_document
from imss.core.models import Bundle, Material, Settings
from imss.workflow import build, scan


class FormFieldsTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.signature = self.root / 'signature.png'
        Image.new('RGB', (400, 100), 'black').save(self.signature)
        self.materials = []
        for index, kind in enumerate(('电子发票', '购买记录')):
            path = self.root / f'material{index}.png'
            Image.new('RGB', (120, 100), ('white', 'gray')[index]).save(path)
            self.materials.append(Material(path.name, str(index), '.png', kind=kind, item='打印纸',
                                           amount='35.80', archived=str(path)))

    def tearDown(self):
        self.temporary.cleanup()

    def test_signature_appears_in_buyer_cell_and_promise_bytes_remain(self):
        destination = self.root / 'form.docx'
        build_document(Bundle('test', '打印纸', '35.80', self.materials), '02', None,
                       destination, Settings(), details={'signature_image': str(self.signature),
                                                        'student_id': 'TEST-ID', 'phone': 'TEST-PHONE'})
        document = Document(destination)
        declaration = document.tables[-1]
        buyer = declaration.cell(0, 0)
        buyer_ids = buyer._tc.xpath('.//a:blip/@r:embed')
        self.assertEqual(len(buyer_ids), 1)
        self.assertEqual(document.part.related_parts[buyer_ids[0]].blob, self.signature.read_bytes())
        promise_ids = declaration.cell(1, 3)._tc.xpath('.//a:blip/@r:embed')
        self.assertEqual(len(promise_ids), 1)
        promise = Path(__file__).resolve().parents[1] / 'imss/resources/templates/承诺不重复报销.png'
        self.assertEqual(document.part.related_parts[promise_ids[0]].blob, promise.read_bytes())
        text = '\n'.join(c.text for r in declaration.rows for c in r.cells)
        self.assertIn('TEST-ID', text)
        self.assertIn('TEST-PHONE', text)
        self.assertNotIn('{{', text)

    def test_travel_names_are_separate_from_signature_and_promise(self):
        for material in self.materials:
            material.category = '行程'
        destination = self.root / 'travel.docx'
        build_document(Bundle('test', '打车费用', '35.80', self.materials), '01', None,
                       destination, Settings(), details={'passenger_name': '张三',
                                                        'passenger_name_2': '李四',
                                                        'passenger_name_3': '王五',
                                                        'student_id': 'TEST-ID',
                                                        'signature_image': str(self.signature)})
        document = Document(destination)
        declaration = document.tables[-1]
        for index, name in enumerate(('张三', '李四', '王五')):
            self.assertIn(f'乘车人{index + 1}姓名：{name}', declaration.cell(index, 0).text)
            self.assertIn('签名：', declaration.cell(index, 0).text)
        self.assertIn('TEST-ID', declaration.cell(0, 3).text)
        signature_id = declaration.cell(0, 0)._tc.xpath('.//a:blip/@r:embed')[0]
        self.assertEqual(document.part.related_parts[signature_id].blob, self.signature.read_bytes())
        promise_id = declaration.cell(4, 3)._tc.xpath('.//a:blip/@r:embed')[0]
        promise = Path(__file__).resolve().parents[1] / 'imss/resources/templates/承诺不重复报销.png'
        self.assertEqual(document.part.related_parts[promise_id].blob, promise.read_bytes())
        self.assertNotIn('{{', '\n'.join(c.text for r in declaration.rows for c in r.cells))

    def test_travel_names_can_be_left_blank_without_signature(self):
        destination = self.root / 'travel.docx'
        build_document(Bundle('test', '打车费用', '35.80', self.materials), '01', None,
                       destination, Settings(), details={'category': '行程'})
        declaration = Document(destination).tables[-1]
        for index in range(3):
            cell = declaration.cell(index, 0)
            self.assertIn(f'乘车人{index + 1}姓名：', cell.text)
            self.assertNotIn('{{', cell.text)
            self.assertFalse(cell._tc.xpath('.//a:blip'))

    def make_scan(self, travel=False):
        source = self.root / 'input'
        source.mkdir()
        for index in range(3 if travel else 2):
            Image.new('RGB', (120 + index, 100), 'white').save(source / f'{index}.png')
        texts = {'0.png': '电子发票\n物品名称：打印纸\n（小写）￥35.80',
                 '1.png': '订单编号：123456\n物品名称：打印纸\n实付款：35.80'}
        if travel:
            texts = {'0.png': '电子发票\n发票号码：12345678901234567890\n物品名称：打车费用\n（小写）￥35.80',
                     '1.png': '订单编号：123456\n物品名称：打车费用\n实付款：35.80',
                     '2.png': '高德地图打车行程单\n合计35.80元'}
        with patch('imss.workflow.extract', side_effect=lambda path, ext, settings: (texts[path.name], [])):
            return scan(source, None, Settings(workers=1))

    def test_group_person_fields_override_main_defaults_and_resolve_relative_images(self):
        manifest = self.make_scan()
        review = manifest.parent / '核对清单.xlsx'
        details = self.root / 'details.json'
        details.write_text(json.dumps({'0.png': {'phone': 'GROUP-PHONE', 'signature_image': 'signature.png'}}), encoding='utf-8')
        from main import main
        with redirect_stdout(io.StringIO()):
            status = main(['build', str(manifest), '--review', str(review), '--details', str(details),
                           '--skip-pdf', '--phone', 'CLI-PHONE'],
                          person_defaults={'student_id': 'DEFAULT-ID', 'phone': 'DEFAULT-PHONE'})
        self.assertEqual(status, 0)
        report = json.loads(manifest.read_text(encoding='utf-8'))['report']
        self.assertFalse(report['pending'])
        document = Document(Path(report['output']) / report['generated'][0]['word'])
        text = '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
        self.assertIn('DEFAULT-ID', text)
        self.assertIn('GROUP-PHONE', text)
        self.assertNotIn('DEFAULT-PHONE', text)
        self.assertNotIn('CLI-PHONE', text)
        self.assertEqual(len(document.tables[-1].cell(0, 0)._tc.xpath('.//a:blip')), 1)

    def test_invalid_signature_is_rejected_before_archiving(self):
        manifest = self.make_scan()
        with self.assertRaisesRegex(ValueError, '签字图片不存在'):
            build(manifest, skip_pdf=True, person_defaults={'signature_image': str(self.root / 'missing.png')})
        self.assertFalse(any((self.root / '报销材料').iterdir()))

    def test_visible_auto_group_details_and_invoice_override(self):
        manifest = self.make_scan(travel=True)
        data = json.loads(manifest.read_text(encoding='utf-8'))
        invoice = next(m for m in data['materials'] if m['kind'] == '电子发票')
        group = 'auto-' + invoice['digest'][:16]
        self.assertTrue(group.startswith('auto-'))
        details = self.root / 'detail.json'
        fields = {group: {'passenger_name': '分组姓名', 'passenger_name_2': '同行人',
                          'purpose': '现场布置往返'}}
        details.write_text(json.dumps(fields, ensure_ascii=False), encoding='utf-8-sig')
        report = build(manifest, skip_pdf=True, details_path=details,
                       person_defaults={'passenger_name': '默认姓名'})
        self.assertFalse(report['pending'])
        def text_in_report(report):
            document = Document(Path(report['output']) / report['generated'][0]['word'])
            return '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
        text = text_in_report(report)
        self.assertIn('乘车人1姓名：分组姓名', text)
        self.assertIn('乘车人2姓名：同行人', text)
        self.assertIn('现场布置往返', text)
        self.assertNotIn('默认姓名', text)
        fields['12345678901234567890'] = {'passenger_name': '发票姓名'}
        details.write_text(json.dumps(fields, ensure_ascii=False), encoding='utf-8')
        report = build(manifest, skip_pdf=True, details_path=details,
                       person_defaults={'passenger_name': '默认姓名'})
        self.assertFalse(report['pending'])
        text = text_in_report(report)
        self.assertIn('乘车人1姓名：发票姓名', text)
        self.assertIn('乘车人2姓名：同行人', text)
        self.assertNotIn('分组姓名', text)

    def test_cli_default_passenger_name(self):
        manifest = self.make_scan(travel=True)
        from main import main
        with redirect_stdout(io.StringIO()):
            status = main(['build', str(manifest), '--passenger-name', '张三', '--skip-pdf'])
        self.assertEqual(status, 0)
        report = json.loads(manifest.read_text(encoding='utf-8'))['report']
        document = Document(Path(report['output']) / report['generated'][0]['word'])
        self.assertIn('乘车人1姓名：张三', document.tables[-1].cell(0, 0).text)

    def test_passenger_names_must_be_strings_before_archiving(self):
        manifest = self.make_scan()
        for field in ('passenger_name', 'passenger_name_2', 'passenger_name_3'):
            with self.subTest(field=field):
                details = self.root / 'detail.json'
                details.write_text(json.dumps({'test': {field: ['张三']}}), encoding='utf-8')
                with self.assertRaisesRegex(ValueError, field + ' 必须是字符串'):
                    build(manifest, skip_pdf=True, details_path=details)
                self.assertFalse(any((self.root / '报销材料').iterdir()))
