import csv
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from PIL import Image, ImageDraw
from docx import Document

from imss.output.documents import build_document, replace_fields
from imss.recognition.extract import detect_extension, extract_pdf, ocr_image
from imss.core.matching import group_materials
from imss.storage import digest_file
from imss.output.review import apply_review, write_review
from imss.core.models import Bundle, Material, Settings, money, safe_name
from imss.workflow import build, scan
from imss.recognition.rules import recognize
from imss.output.workspace import default_output, default_records


def record(source, kind, item='打印纸', amount='35.80', **kwargs):
    return Material(source=source, digest=source, extension='.png', kind=kind, item=item, amount=amount, **kwargs)


def text_pdf(path):
    # A tiny real PDF with a long text layer, independent of OCR/Office/reportlab.
    stream = b'BT /F1 12 Tf 40 180 Td (A readable invoice text layer 1234567890) Tj ET'
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 220] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    result = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for index, content in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(str(index).encode() + b' 0 obj\n' + content + b'\nendobj\n')
    xref = len(result)
    result.extend(b'xref\n0 6\n0000000000 65535 f \n')
    for offset in offsets[1:]:
        result.extend(f'{offset:010d} 00000 n \n'.encode())
    result.extend(f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode())
    path.write_bytes(result)


class RulesTests(unittest.TestCase):
    def test_decimal_money(self):
        self.assertEqual(money('1,439'), '1439.00')
        for invalid in ('NaN', 'Infinity', '0', '-1', '1.234', '1000000000'):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                money(invalid)

    def test_invoice_gross_total(self):
        material = record('i', '未知', item='', amount='')
        material.text = '电子发票\n*办公用品*打印纸\n金额 31.68 税额4.12\n价税合计（大写）叁拾伍圆捌角（小写）￥35.80'
        recognize(material)
        self.assertEqual((material.kind, material.item, material.amount), ('电子发票', '打印纸', '35.80'))

    def test_ocr_spacing(self):
        material = record('p', '未知', item='', amount='')
        material.text = '交易 成 功\n商品 名称：打印纸\n实 付款 HY ¥35.80'
        recognize(material)
        self.assertEqual(material.amount, '35.80')

    def test_title_beats_invoice_mention(self):
        material = record('d', '未知')
        material.text = '承诺书\n本人承诺发票号码对应支出不重复报销'
        recognize(material)
        self.assertEqual(material.kind, '承诺书')

    def test_itinerary_distinct_from_invoice(self):
        material = record('t', '未知')
        material.text = '高德地图打车行程单\n共计2单行程，合计52.92元'
        recognize(material)
        self.assertEqual((material.kind, material.amount, material.category), ('行程单', '52.92', '行程'))

    def test_news_title_ignores_site_navigation(self):
        material = record('n', '未知')
        material.text = '站内搜索\n网站首页 学工新闻\n某研究院毕业典礼暨毕业晚会举行\n发布者：编辑 发布时间：2026-06-25 浏览次数：19'
        recognize(material)
        self.assertEqual(material.kind, '新闻稿')
        self.assertEqual(material.title, '某研究院毕业典礼暨毕业晚会举行')

    def test_multiple_goods_require_review(self):
        material = record('i', '未知')
        material.text = '电子发票\n*商品*胶带 1\n*商品*气球 2\n（小写）￥35.80'
        recognize(material)
        self.assertEqual(material.item, '')
        self.assertTrue(any('多个商品' in issue for issue in material.issues))

    def test_ambiguous_invoices_not_merged(self):
        bundles, pending = group_materials([record('i1', '电子发票'), record('i2', '电子发票'), record('p', '购买记录')])
        self.assertEqual(len(bundles), 2)
        self.assertEqual([m.source for m in pending], ['p'])

    def test_explicit_orders_must_not_conflict(self):
        _, pending = group_materials([record('i', '电子发票', order_ids=['111111']), record('p', '购买记录', order_ids=['222222'])])
        self.assertEqual(len(pending), 1)

    def test_fare_total_cannot_steal_another_groups_subfare(self):
        invoice_one = record('i1', '电子发票', item='打车费用', amount='19.80', category='行程')
        invoice_two = record('i2', '电子发票', item='打车费用', amount='45.00', category='行程')
        itinerary_one = record('t1', '行程单', item='打车费用', amount='19.80', category='行程', related_amounts=['19.80'])
        itinerary_two = record('t2', '行程单', item='打车费用', amount='45.00', category='行程', related_amounts=['19.80', '25.20'])
        payment = record('p', '购买记录', item='打车费用', amount='19.80', category='行程')
        bundles, pending = group_materials([invoice_one, invoice_two, itinerary_one, itinerary_two, payment])
        self.assertIn(payment, pending)
        self.assertTrue(all(payment not in bundle.materials for bundle in bundles))

    def test_manual_group_cannot_hide_amount_conflict(self):
        bundles, _ = group_materials([record('i', '电子发票', group='a'), record('p', '购买记录', amount='35.81', group='a')])
        self.assertTrue(any('金额' in issue for issue in bundles[0].issues))

    def test_manual_split_payments_sum_to_invoice(self):
        bundles, pending = group_materials([record('i', '电子发票', amount='480.00', group='glow'),
                                           record('p1', '购买记录', amount='80.00', group='glow'),
                                           record('p2', '购买记录', amount='400.00', group='glow')])
        self.assertFalse(pending)
        self.assertFalse(bundles[0].issues)
        self.assertEqual(bundles[0].amount, '480.00')

    def test_shared_receipt_sheet_is_used_by_both_groups(self):
        sheet = record('sheet', '签领表', item='', amount='', group='fan|cup')
        bundles, pending = group_materials([record('if', '电子发票', group='fan'), record('pf', '购买记录', group='fan'),
                                           record('ic', '电子发票', group='cup'), record('pc', '购买记录', group='cup'), sheet])
        self.assertFalse(pending)
        self.assertTrue(all(not bundle.issues for bundle in bundles))
        self.assertTrue(all(sheet in bundle.materials for bundle in bundles))

    def test_manual_group_order_is_independent_of_random_filenames(self):
        materials = []
        for name, group in [('z', '2'), ('a', '10'), ('c', '24.2'), ('b', '24.1')]:
            materials.extend([record(name + 'i', '电子发票', group=group),
                              record(name + 'p', '购买记录', group=group)])
        bundles, pending = group_materials(materials)
        self.assertFalse(pending)
        self.assertEqual([b.key for b in bundles], ['manual:2', 'manual:10', 'manual:24.1', 'manual:24.2'])

    def test_shared_order_inherits_blank_declaration_fields(self):
        declaration = record('d', '承诺书', item='', amount='', order_ids=['111111'])
        declaration.issues = ['未找到明确的物品名称', '未找到明确的含税合计或实付金额']
        bundles, pending = group_materials([record('i', '电子发票', order_ids=['111111']), record('p', '购买记录'), declaration])
        self.assertFalse(pending)
        self.assertFalse(bundles[0].issues)
        self.assertEqual((declaration.item, declaration.amount), ('打印纸', '35.80'))

    def test_safe_names(self):
        self.assertEqual(safe_name('a/b:c?'), 'a_b_c_')
        self.assertEqual(safe_name('CON'), '_CON')
        self.assertLessEqual(len(safe_name('气' * 200)), 64)


class FileTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def test_temp_image_sniffing(self):
        path = self.root / 'a.temp'
        Image.new('RGB', (200, 150), 'white').save(path, format='PNG')
        self.assertEqual(detect_extension(path), '.png')

    def test_real_pdf_extraction_without_ocr(self):
        path = self.root / 'a.tmp'
        text_pdf(path)
        with patch('imss.recognition.extract.ocr_image', side_effect=AssertionError('should not OCR')):
            text, issues = extract_pdf(path, Settings())
        self.assertIn('readable invoice', text)
        self.assertFalse(issues)
        self.assertEqual(detect_extension(path), '.pdf')

    def test_split_run_template_fields(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('金额 {{am').bold = True
        paragraph.add_run('ount}} 元；{{amount}}')
        replace_fields(paragraph, {'amount': '35.80'})
        self.assertEqual(paragraph.text, '金额 35.80 元；35.80')
        self.assertTrue(paragraph.runs[0].bold)

    def test_review_requires_explicit_confirmation(self):
        material = record('a', '电子发票')
        path = self.root / 'review.csv'
        write_review(path, [material])
        with path.open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.DictReader(stream))
        rows[0]['item'] = '气球'
        def save():
            with path.open('w', encoding='utf-8-sig', newline='') as stream:
                writer = csv.DictWriter(stream, fieldnames=rows[0].keys())
                writer.writeheader()
                writer.writerows(rows)
        save()
        apply_review(path, [material])
        self.assertEqual(material.item, '打印纸')
        rows[0]['confirmed'] = '1'
        save()
        apply_review(path, [material])
        self.assertEqual(material.item, '气球')

    def test_run_applies_review_and_business_details(self):
        from contextlib import redirect_stdout
        from main import main

        source = self.root / 'input'
        source.mkdir()
        rows = [record('i.png', '电子发票', item='彩排买水', amount='20.00', group='01'),
                record('p.png', '购买记录', item='彩排买水', amount='20.00', group='01')]
        for index, material in enumerate(rows):
            Image.new('RGB', (100 + index, 90), 'white').save(source / material.source)
        review = self.root / 'review.csv'
        write_review(review, rows)
        with review.open(encoding='utf-8-sig', newline='') as stream:
            reviewed = list(csv.DictReader(stream))
        with review.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=reviewed[0].keys())
            writer.writeheader()
            writer.writerows(dict(row, confirmed='1') for row in reviewed)
        details = self.root / 'details.json'
        details.write_text(json.dumps({'01': {'purpose': '彩排用水', 'invoice_number': '12345678901234567890',
                                               'is_prize': False, 'amount_equal': True}}), encoding='utf-8')
        output = default_output(source)
        with patch('imss.workflow.extract', return_value=('无法识别的材料', [])), redirect_stdout(io.StringIO()):
            result = main(['run', str(source), '--review', str(review),
                           '--details', str(details), '--skip-pdf'])
        self.assertEqual(result, 0)
        report = json.loads((default_records(output) / '核对清单.json').read_text(encoding='utf-8'))['report']
        self.assertEqual(report['review_file'], str(review.resolve()))
        self.assertEqual(report['generated'][0]['group'], 'manual:01')
        self.assertEqual(set(report['generated'][0]['sources']), {'i.png', 'p.png'})
        document = Document(output / report['generated'][0]['word'])
        text = '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
        self.assertIn('彩排用水', text)
        self.assertIn('12345678901234567890', text)

    def test_tsv_without_config_files(self):
        import subprocess
        tsv = b'level\tpage_num\tblock_num\tpar_num\tline_num\tconf\ttext\n5\t1\t1\t1\t1\t96\tHello\n'
        with patch('imss.recognition.extract.tessdata_path', return_value=Path('models-only')), patch('imss.recognition.extract.tesseract_path', return_value='tesseract'), patch('imss.recognition.extract.run_process', return_value=subprocess.CompletedProcess([], 0, tsv, b'')) as process:
            text, issues = ocr_image(Image.new('RGB', (100, 100)), Settings(tessdata='models-only'))
        self.assertEqual(text, 'Hello')
        self.assertFalse(issues)
        self.assertIn('tessedit_create_tsv=1', process.call_args.args[0])

    def test_templates_retain_only_the_authorized_promise_image(self):
        root = Path(__file__).resolve().parents[1] / 'imss' / 'resources' / 'templates'
        promise = (root / '承诺不重复报销.png').read_bytes()
        for path in root.glob('*.docx'):
            with self.subTest(template=path.name), zipfile.ZipFile(path) as archive:
                images = [name for name in archive.namelist() if name.startswith('word/media/')]
                self.assertEqual(len(images), 1)
                self.assertEqual(archive.read(images[0]), promise)
                xml = archive.read('word/document.xml').decode('utf-8')
                self.assertNotRegex(xml, r'\d{11}|SA\d{8}|\d{20}')

    def test_numbering_archive_and_originals_unchanged(self):
        source = self.root / 'input'
        source.mkdir()
        texts = {'i.tmp': '电子发票\n物品名称：打印纸\n（小写）￥35.80',
                 'p.png': '订单编号：123456\n物品名称：打印纸\n实付款：35.80',
                 'a.png': '项目审批表\n测试项目', 'n.png': '新闻稿\n测试活动', 's.png': '情况说明\n测试支出'}
        for index, name in enumerate(texts):
            Image.new('RGB', (100 + index, 90), 'white').save(source / name, format='PNG')
        hashes = {name: digest_file(source / name) for name in texts}
        with patch('imss.workflow.extract', side_effect=lambda path, ext, settings: (texts[path.name], [])):
            manifest = scan(source, None, Settings(workers=1))
        report = build(manifest, skip_pdf=True)
        self.assertFalse(report['pending'])
        self.assertEqual([value['sequence'] for value in report['generated']], ['00', '00', '01', '02'])
        output = Path(report['output'])
        self.assertTrue((output / '01.打印纸_35.80.docx').is_file())
        self.assertEqual({p.name for p in (output / '原始材料').glob('*元-*')},
                         {'01.打印纸-35.80元-电子发票.png', '01.打印纸-35.80元-购买记录.png'})
        self.assertTrue(all(not m['archived'] for m in report['materials'] if m['kind'] in ('审查表', '新闻稿', '情况说明')))
        self.assertFalse(any(p.suffix in ('.json', '.csv') for p in output.iterdir()))
        self.assertTrue((output / '00.项目审批表审查表.pdf').is_file())
        news = next(m['pdf'] for m in report['generated'] if m['kind'] == '新闻稿')
        self.assertEqual(Path(news).parent, Path('.'))
        self.assertTrue((output / news).is_file())
        self.assertEqual({m['pdf'] for m in report['generated'] if m['kind'] == '审查表'},
                         {'00.项目审批表审查表.pdf'})
        self.assertEqual(hashes, {name: digest_file(source / name) for name in texts})
        # 重复生成读取同一张核对表，原件与输出布局保持不变。
        second = build(manifest, skip_pdf=True)
        self.assertFalse(second['pending'])
        self.assertEqual({v['word'] for v in second['generated'] if v.get('word')}, {'01.打印纸_35.80.docx'})
        self.assertEqual({p.name for p in output.iterdir() if p.is_dir()}, {'原始材料'})

    def test_changes_after_scan_detected(self):
        source = self.root / 'input'
        source.mkdir()
        Image.new('RGB', (50, 50)).save(source / 'a.png')
        with patch('imss.workflow.extract', return_value=('项目审批表', [])):
            manifest = scan(source, None, Settings(workers=1))
        (source / 'a.png').write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError, '发生变化'):
            build(manifest)

    def test_cli_defaults_to_clean_sibling_output_and_preserves_originals(self):
        from contextlib import redirect_stdout
        from main import main

        source = self.root / 'input'
        source.mkdir()
        texts = {'i.png': '电子发票\n物品名称：打印纸\n（小写）￥35.80',
                 'p.png': '订单编号：123456\n物品名称：打印纸\n实付款：35.80'}
        for index, name in enumerate(texts):
            Image.new('RGB', (100 + index, 90), 'white').save(source / name)
        hashes = {name: digest_file(source / name) for name in texts}
        with patch('imss.workflow.extract', side_effect=lambda path, ext, settings: (texts[path.name], [])), redirect_stdout(io.StringIO()):
            self.assertEqual(main(['run', str(source), '--skip-pdf']), 0)
            output = default_output(source)
            report_path = default_records(output) / '核对清单.json'
            report = json.loads(report_path.read_text(encoding='utf-8'))['report']
            self.assertEqual(Path(report['output']), output)
            self.assertTrue((output / '01.打印纸_35.80.docx').is_file())
            self.assertEqual(len(list((output / '原始材料').glob('*元-*'))), 2)
            self.assertEqual({p.name for p in output.iterdir() if p.is_dir()}, {'原始材料'})
            self.assertEqual({p.name for p in output.parent.iterdir()}, {'input', '报销材料', '核对清单'})
            self.assertEqual({p.name for p in report_path.parent.iterdir()}, {'核对清单.json', '核对清单.xlsx'})
            self.assertEqual(main(['scan', str(source)]), 0)
            self.assertFalse(any(output.iterdir()))
            rescanned = json.loads(report_path.read_text(encoding='utf-8'))
            self.assertNotIn('report', rescanned)
            self.assertEqual(len(rescanned['materials']), 2)
        self.assertEqual(hashes, {name: digest_file(source / name) for name in texts})

    def test_in_place_scan_preserves_existing_reserved_materials(self):
        existing = self.root / '原始材料'
        existing.mkdir()
        path = existing / 'a.png'
        Image.new('RGB', (50, 50), 'white').save(path)
        original = digest_file(path)
        with self.assertRaisesRegex(ValueError, '输入和输出不能同目录'):
            scan(self.root, self.root, Settings())
        self.assertEqual(digest_file(path), original)
        self.assertFalse((self.root / '扫描结果.json').exists())


if __name__ == '__main__':
    unittest.main()
