"""多证据、坐标提取和批量核对的行为回归，不依赖私人材料标签。"""
from copy import deepcopy
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook

from imss.core.candidates import pairing_candidates
from imss.core.matching import group_materials
from imss.core.models import Material, Settings
from imss.output.review import apply_review, read_review
from tests.legacy_review_fixture import write_review
from imss.recognition.extract import Extraction, ocr_image
from imss.recognition.rules import recognize
from imss.recognition.refine import refine_missing_amount
from imss.workflow import build, recheck, scan


def record(source, kind='电子发票', item='透明胶带', amount='22.50', **values):
    return Material(source, source, '.png', kind=kind, item=item, amount=amount, category='物品', **values)


def block(text, x, y, width=.15, height=.03, page=1):
    return dict(text=text, box=[x, y, x + width, y + height], confidence=93, page=page, engine='tesseract')


class MultiEvidenceTests(unittest.TestCase):
    def test_marketing_name_proposes_without_changing_or_auto_approving(self):
        materials = [record('invoice'), record('order', '购买记录', item='大卷透明胶带批发宽款')]
        before = deepcopy(materials)
        proposals = pairing_candidates(materials)
        self.assertEqual(materials, before)
        self.assertEqual(len(proposals), 1)
        self.assertIn('包含', proposals[0]['reasons'])
        self.assertIn('名称不同', proposals[0]['risks'])
        self.assertTrue(group_materials(materials)[1])

    def test_rank_uses_independent_signals_and_retains_competitors(self):
        proposals = pairing_candidates([record('i1', merchant='甲商店'), record('i2', item='水杯'),
            record('o', '购买记录', item='大卷透明胶带', merchant='甲商店')])
        self.assertEqual(len(proposals), 2)
        self.assertEqual(proposals[0]['invoice'], 'i1')
        self.assertIn('商家一致', proposals[0]['reasons'])
        self.assertTrue(all('竞争' in p['status'] for p in proposals))

    def test_conflicting_identifiers_and_amounts_are_excluded(self):
        for values in (dict(amount='21.00'), dict(order_ids=['different']), dict(invoice_ids=['different'])):
            invoice = record('i', order_ids=['same'], invoice_ids=['same'])
            order = record('o', '购买记录', item='大卷透明胶带', **values)
            self.assertFalse(pairing_candidates([invoice, order]))

    def test_inferred_item_is_never_independent_name_evidence(self):
        materials = [record('i'), record('o', '购买记录', item='')]
        group_materials(materials)
        proposal = pairing_candidates(materials)[0]
        self.assertNotIn('名称一致', proposal['reasons'])
        self.assertEqual(proposal['order_item'], '')
        self.assertIn('只有金额支持', proposal['risks'])

    def test_identifier_pair_with_name_mismatch_can_be_confirmed_as_group(self):
        materials = [record('i', order_ids=['same']),
                     record('o', '购买记录', item='大卷透明胶带', order_ids=['same'])]
        self.assertEqual(len(pairing_candidates(materials)), 1)

    def test_ocr_keeps_word_boxes_scaled_to_original_page(self):
        tsv = ('text\tconf\tline_num\tpage_num\tblock_num\tpar_num\tleft\ttop\twidth\theight\n'
               '实付款\t96\t1\t1\t1\t1\t40\t80\t100\t30\n'
               '22.50\t93\t1\t1\t1\t1\t200\t80\t100\t30\n').encode()
        with patch('imss.recognition.extract.tesseract_path', return_value='tesseract'), patch(
                'imss.recognition.extract.run_process', return_value=subprocess.CompletedProcess([], 0, tsv, b'')):
            result = ocr_image(Image.new('RGB', (250, 500)), Settings(ocr_min_chars=1))
        self.assertEqual(result.layout[0]['box'], [.04, .04, .14, .055])
        self.assertEqual(result.layout[1]['confidence'], 93)
        self.assertEqual(len(result), 2)  # 原来的解包接口不变。

    def test_spatial_label_value_repairs_fragmented_order_total(self):
        material = Material('order', 'digest', '.png', text='交易成功\n实付款\n优惠\n22.50\n2.00',
            layout=[block('实付款', .1, .1), block('22.50', .7, .1),
                    block('优惠', .1, .2), block('2.00', .7, .2)])
        recognize(material)
        self.assertEqual(material.amount, '22.50')
        self.assertTrue(material.field_evidence['amount']['regions'])
        self.assertEqual(material.field_evidence['amount']['regions'][0]['page'], 1)

    def test_different_pages_are_never_merged_into_one_payment_line(self):
        material = Material('order', 'digest', '.png', text='交易成功\n实付款',
            layout=[block('实付款', .1, .1), block('22.50', .7, .1, page=2)])
        recognize(material)
        self.assertFalse(material.amount)
        # 缺失字段可以提示；不能制造跨页同一行证据。
        for entry in material.field_evidence.values():
            self.assertTrue(all(region['text'] != '实付款 22.50' for region in entry['regions']))

    def test_spatial_payment_label_may_have_spaces_between_chinese_characters(self):
        material = Material('order', 'digest', '.png', text='交易成功\n实付款共减 Y4.8\nY45.20',
            layout=[block('实 付 款', .1, .1, width=.2), block('共减Y4.8', .4, .1), block('Y45.20', .7, .1)])
        recognize(material)
        self.assertEqual(material.amount, '45.20')

    def test_local_numeric_ocr_keeps_warning_and_never_guesses_decimal(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'bill.png'
            Image.new('RGB', (300, 700), 'white').save(path)
            for text, expected in (('-45.20', '45.20'), ('4520', '')):
                material = record('o', '购买记录', amount='', layout=[block('当前状态 支付成功', .1, .4)])
                tsv = ('text\tconf\tleft\ttop\twidth\theight\n' + text + '\t95\t30\t30\t150\t30\n').encode()
                with patch('imss.recognition.refine.tesseract_path', return_value='tesseract'), patch(
                        'imss.recognition.refine.run_process', return_value=subprocess.CompletedProcess([], 0, tsv, b'')):
                    refine_missing_amount(path, material, Settings())
                self.assertEqual(material.amount, expected)
                if expected:
                    self.assertTrue(any('确认后才可生成' in issue for issue in material.issues))
                    self.assertTrue(material.field_evidence['amount']['regions'])

    def test_local_ocr_cannot_resolve_explicit_amount_conflict(self):
        material = record('o', '购买记录', amount='', issues=['实付金额冲突: 4.80, 45.20'],
                          layout=[block('当前状态 支付成功', .1, .4)])
        with patch('imss.recognition.refine.run_process') as process:
            refine_missing_amount(Path('not-read.png'), material, Settings())
        process.assert_not_called()


class BatchConfirmationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / '核对清单.xlsx'

    def confirm(self, indices=(5,), edit=None):
        workbook = load_workbook(self.path)
        sheet = workbook['候选配对']
        for index in indices:
            sheet.cell(index, 1, '1')
        if edit:
            edit(workbook)
        workbook.save(self.path)
        workbook.close()

    def test_batch_confirmation_canonicalizes_and_passes_validation(self):
        materials = [record('i'), record('o', '购买记录', item='大卷透明胶带')]
        write_review(self.path, materials, self.root)
        self.confirm()
        apply_review(self.path, materials)
        bundles, pending = group_materials(materials)
        self.assertFalse(pending)
        self.assertFalse(bundles[0].issues)
        self.assertTrue(all(m.confirmed and m.item == '透明胶带' for m in materials))

    def test_competing_confirmations_are_rejected(self):
        materials = [record('i'), record('o1', '购买记录', item='大卷透明胶带'),
                     record('o2', '购买记录', item='宽透明胶带')]
        write_review(self.path, materials, self.root)
        self.confirm((5, 6))
        with self.assertRaisesRegex(ValueError, '争用'):
            apply_review(self.path, materials)

    def test_multiple_independent_candidates_can_be_confirmed_together(self):
        materials = [record('i1'), record('o1', '购买记录', item='大卷透明胶带'),
                     record('i2', item='纸杯', amount='45.20'),
                     record('o2', '购买记录', item='一次性纸杯批发', amount='45.20')]
        write_review(self.path, materials, self.root)
        self.confirm((5, 6))
        apply_review(self.path, materials)
        bundles, pending = group_materials(materials)
        self.assertEqual(len(bundles), 2)
        self.assertFalse(pending)
        self.assertTrue(all(not b.issues for b in bundles))

    def test_locked_review_preserves_original_and_removes_temporary_file(self):
        write_review(self.path, [record('i')])
        before = self.path.read_bytes()
        with patch('pathlib.Path.replace', side_effect=PermissionError('locked')):
            with self.assertRaisesRegex(PermissionError, '关闭 Excel'):
                write_review(self.path, [record('i')])
        self.assertEqual(before, self.path.read_bytes())
        self.assertFalse(self.path.with_suffix('.partial.xlsx').exists())

    def test_altered_or_stale_candidate_is_rejected(self):
        materials = [record('i'), record('o', '购买记录', item='大卷透明胶带')]
        write_review(self.path, materials, self.root)
        self.confirm(edit=lambda w: setattr(w['候选配对']['J5'], 'value', '0.01'))
        with self.assertRaisesRegex(ValueError, '候选内容已修改'):
            apply_review(self.path, deepcopy(materials))
        write_review(self.path, materials, self.root)
        self.confirm()
        materials[1].digest = 'new-content'
        with self.assertRaisesRegex(ValueError, '失效'):
            apply_review(self.path, materials)

    def test_manual_edits_cannot_be_silently_overridden_by_batch(self):
        materials = [record('i'), record('o', '购买记录', item='大卷透明胶带')]
        write_review(self.path, materials, self.root)
        self.confirm(edit=lambda w: setattr(w['核对清单']['D5'], 'value', '手改名称'))
        with self.assertRaisesRegex(ValueError, '两张表'):
            apply_review(self.path, materials)

    def test_confirmation_does_not_bypass_required_travel_attachment(self):
        materials = [record('i'), record('o', '购买记录', item='大卷透明胶带')]
        for m in materials:
            m.category = '行程'
        write_review(self.path, materials, self.root)
        self.confirm()
        apply_review(self.path, materials)
        self.assertIn('行程报销缺少行程单', group_materials(materials)[0][0].issues)

    def test_legacy_workbook_without_candidates_or_evidence_still_reads(self):
        materials = [record('i')]
        write_review(self.path, materials)
        workbook = load_workbook(self.path)
        del workbook['候选配对']
        workbook.active.delete_cols(11, 3)
        workbook.save(self.path)
        workbook.close()
        self.assertEqual(len(read_review(self.path)), 1)
        apply_review(self.path, materials)

    def test_build_repeats_batch_choice_and_recheck_keeps_it(self):
        source = self.root / 'input'
        source.mkdir()
        for index, name in enumerate(('i.png', 'o.png')):
            Image.new('RGB', (100 + index, 100), 'white').save(source / name)
        texts = {'i.png': '电子发票\n物品名称：透明胶带\n（小写）￥22.50',
                 'o.png': '交易成功\n物品名称：大卷透明胶带\n实付款：22.50'}
        with patch('imss.workflow.extract', side_effect=lambda path, *_: Extraction(texts[path.name], [])):
            manifest = scan(source, None, Settings(workers=1))
        self.path = manifest.with_suffix('.xlsx')
        write_review(self.path, [Material(**m) for m in json.loads(manifest.read_text(encoding='utf-8'))['materials']], source)
        self.confirm()
        first = build(manifest, skip_pdf=True)
        second = build(manifest, skip_pdf=True)
        self.assertEqual(len(first['generated']), 1)
        self.assertFalse(second['pending'])
        with patch('imss.workflow.extract') as extraction:
            result = recheck(manifest)
        extraction.assert_not_called()
        self.assertEqual(result['confirmed'], 2)
        third = build(manifest, skip_pdf=True)
        self.assertFalse(third['pending'])

    def test_recheck_keeps_unconfirmed_edits_and_refreshes_layout(self):
        source = self.root / 'input'
        source.mkdir()
        Image.new('RGB', (100, 100)).save(source / 'i.png')
        with patch('imss.workflow.extract', return_value=('交易成功\n实付款：22.50', [])):
            manifest = scan(source, None, Settings(workers=1))
        self.path = manifest.with_suffix('.xlsx')
        workbook = load_workbook(self.path)
        workbook.active['F5'] = '人工待确认名称'
        workbook.save(self.path)
        workbook.close()
        with patch('imss.workflow.extract', return_value=Extraction('交易成功\n实付款：22.50', [], [block('实付款22.50', .1, .1)])):
            recheck(manifest)
        self.assertEqual(read_review(self.path)[0]['item'], '人工待确认名称')
        self.assertTrue(json.loads(manifest.read_text(encoding='utf-8'))['materials'][0]['layout'])


if __name__ == '__main__':
    unittest.main()
