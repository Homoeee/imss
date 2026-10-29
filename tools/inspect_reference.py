"""Read-only reference inventory. Outputs may contain personal data: keep them local."""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import zipfile

from docx import Document
from lxml import etree
import pypdfium2 as pdfium


def inventory(root: Path):
    result = {"root": str(root.resolve()), "counts": {}, "documents": [], "originals": [], "specials": []}
    raw = root / "原始材料"
    hashes = {}
    for path in sorted(raw.iterdir()):
        if not path.is_file():
            continue
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        hashes.setdefault(digest, []).append(path.name)
        record = {"name": path.name, "extension": path.suffix, "bytes": path.stat().st_size, "sha256": digest}
        if path.suffix.lower() == ".pdf":
            with pdfium.PdfDocument(path) as pdf:
                record["pages"] = len(pdf)
                record["text"] = "\n".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))
        result["originals"].append(record)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
          "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    for path in sorted(root.glob("*.docx")):
        document = Document(path)
        with zipfile.ZipFile(path) as archive:
            xml = etree.fromstring(archive.read("word/document.xml"))
            text = "\n".join("".join(p.xpath(".//w:t/text()", namespaces=ns)) for p in xml.xpath(".//w:p", namespaces=ns))
            media = []
            for name in archive.namelist():
                if name.startswith("word/media/") and not name.endswith("/"):
                    data = archive.read(name)
                    digest = hashlib.sha256(data).hexdigest()
                    media.append({"part": name, "bytes": len(data), "sha256": digest, "matches_original": hashes.get(digest, [])})
            drawings = []
            for drawing in xml.xpath(".//w:drawing", namespaces=ns):
                drawings.append({"xml": etree.tostring(drawing, encoding="unicode")})
            record = {"name": path.name, "text": text, "media": media,
                      "sections": [{"width_mm": s.page_width.mm, "height_mm": s.page_height.mm,
                                    "left_mm": s.left_margin.mm, "right_mm": s.right_margin.mm,
                                    "top_mm": s.top_margin.mm, "bottom_mm": s.bottom_margin.mm} for s in document.sections],
                      "tables": len(document.tables), "drawings": drawings}
        pdf_path = path.with_suffix(".pdf")
        if pdf_path.exists():
            with pdfium.PdfDocument(pdf_path) as pdf:
                record["pdf_pages"] = len(pdf)
                record["pdf_text"] = "\n".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))
        result["documents"].append(record)
    root_docx = {p.stem for p in root.glob("*.docx")}
    for path in sorted(root.glob("*.pdf")):
        if path.stem not in root_docx:
            with pdfium.PdfDocument(path) as pdf:
                result["specials"].append({"name": path.name, "pages": len(pdf),
                                           "text": "\n".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))})
    result["counts"] = {"originals": len(result["originals"]), "original_extensions": dict(Counter(p["extension"] for p in result["originals"])),
                        "word": len(result["documents"]), "specials": len(result["specials"]),
                        "word_pdf_pages": dict(Counter(d.get("pdf_pages", 0) for d in result["documents"]))}
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("reference", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    result = inventory(args.reference)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result["counts"], ensure_ascii=False, indent=2))
    for document in result["documents"]:
        print(document["name"], "tables=", document["tables"], "images=", len(document["media"]), "pages=", document.get("pdf_pages"), "text=", document["text"][:1000].replace("\n", " | "))
